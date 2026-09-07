"""Add current O25 and O100 master tables to the canonical Results TD report."""

from __future__ import annotations

import json
import math
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

SECTION_START = "<!-- RESULTS_TD_O25_O100_CURRENT_MASTER_START -->"
SECTION_END = "<!-- RESULTS_TD_O25_O100_CURRENT_MASTER_END -->"
DOCX_O25_MARKER = "RESULTS TD / O25 CURRENT MASTER TABLE END"
DOCX_O100_MARKER = "RESULTS TD / O100 CURRENT MASTER TABLE END"

SCORE_KEYS = (
    "f_only",
    "g_only",
    "f_plus_g",
    "f_plus_g_first",
    "g_only_f_rollout_mean",
    "f_plus_g_first_q2",
    "state_v_plus_first_q2",
)
SCORE_HEADERS = (
    "F-only",
    "G/C4-only",
    "F+G/F+C4 tail",
    "First-Q alpha=.25",
    "Mean-Q",
    "First-Q2 alpha=.25",
    "C3 State-V+First-Q2 alpha=.10",
)
RESULT_PATHS = {
    "f_only": Path("f_only/results.json"),
    "g_only": Path("g_only/results.json"),
    "f_plus_g": Path("f_plus_g/results.json"),
    "f_plus_g_first": Path("f_plus_g_first/alpha_0p25/results.json"),
    "g_only_f_rollout_mean": Path("g_only_f_rollout_mean/results.json"),
    "f_plus_g_first_q2": Path("f_plus_g_first_q2/alpha_0p25/results.json"),
}


class ProtocolMasterError(ValueError):
    """Raised when evidence or the canonical report has drifted."""


@dataclass(frozen=True)
class ScoreCell:
    count: int
    outcomes: tuple[bool, ...]


@dataclass(frozen=True)
class MethodRow:
    label: str
    training_loss: str
    cells: Mapping[str, ScoreCell | None]


@dataclass(frozen=True)
class ProtocolMaster:
    protocol: str
    rows: tuple[MethodRow, ...]
    selection: tuple[int, ...]


@dataclass(frozen=True)
class MasterEvidence:
    o25: ProtocolMaster
    o100: ProtocolMaster


def _load_json(path: Path) -> Mapping[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(path)
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ProtocolMasterError(f"JSON object required: {path}")
    return payload


def _load_result(path: Path, *, protocol: str, expected_mode: str) -> ScoreCell:
    payload = _load_json(path)
    if str(payload.get("protocol_label", "")).lower() != protocol:
        raise ProtocolMasterError(f"protocol mismatch: {path}")
    if payload.get("score_mode") != expected_mode:
        raise ProtocolMasterError(f"score mode mismatch: {path}")
    if payload.get("smoke") is not False or payload.get("pilot") is not False:
        raise ProtocolMasterError(f"non-formal result in current master: {path}")
    expected_offset = 25 if protocol == "o25" else 100
    expected_budget = 50 if protocol == "o25" else 200
    expected_planning_horizon = 1 if expected_mode == "g_only" else 5
    expected_receding_horizon = (
        1 if protocol == "o100" or expected_mode == "g_only" else 5
    )
    identity = {
        "implementation_version": "v1",
        "method_family": "actor_free_td_lewm_v1",
        "goal_offset": expected_offset,
        "episode_budget": expected_budget,
        "planning_horizon": expected_planning_horizon,
        "receding_horizon": expected_receding_horizon,
    }
    for key, expected in identity.items():
        if payload.get(key) != expected:
            raise ProtocolMasterError(f"{key} mismatch: {path}")
    variants = [
        path.parts[index + 1]
        for index, part in enumerate(path.parts[:-1])
        if part == "v1" and path.parts[index + 1] in {"c", "c3", "c4"}
    ]
    if len(variants) != 1:
        raise ProtocolMasterError(f"cannot identify method variant from path: {path}")
    variant = variants[0]
    if payload.get("variant") != variant:
        raise ProtocolMasterError(f"variant mismatch: {path}")
    if payload.get("method") != f"actor_free_td_lewm_v1_{variant}":
        raise ProtocolMasterError(f"method mismatch: {path}")
    if variant == "c4":
        if payload.get("objective_version") != 1:
            raise ProtocolMasterError(f"C4 objective version mismatch: {path}")
        if (
            payload.get("state_only_g") is not True
            or payload.get("action_enters_g") is not False
        ):
            raise ProtocolMasterError(f"C4 state-only interface mismatch: {path}")
    elif payload.get("objective_version") is not None:
        raise ProtocolMasterError(f"unexpected objective version: {path}")
    metrics = payload.get("metrics")
    if not isinstance(metrics, dict):
        raise ProtocolMasterError(f"metrics object missing: {path}")
    raw_outcomes = metrics.get("episode_successes")
    if (
        not isinstance(raw_outcomes, list)
        or len(raw_outcomes) != 50
        or any(type(value) is not bool for value in raw_outcomes)
    ):
        raise ProtocolMasterError(f"50 Boolean outcomes required: {path}")
    outcomes = tuple(raw_outcomes)
    count = sum(outcomes)
    raw_rate = metrics.get("success_rate")
    if (
        isinstance(raw_rate, bool)
        or not isinstance(raw_rate, (int, float))
        or not math.isfinite(float(raw_rate))
        or abs(float(raw_rate) - 2.0 * count) > 1e-9
    ):
        raise ProtocolMasterError(f"success rate disagrees with outcomes: {path}")
    return ScoreCell(count=count, outcomes=outcomes)


def _selection_signature(path: Path, *, protocol: str) -> tuple[tuple[int, ...], ...]:
    payload = _load_json(path.parent / "episode_selection.json")
    keys = ("episode_indices", "start_steps", "goal_steps", "valid_row_ranks")
    fields: list[tuple[int, ...]] = []
    for key in keys:
        values = payload.get(key)
        if (
            not isinstance(values, list)
            or len(values) != 50
            or any(type(value) is not int for value in values)
        ):
            raise ProtocolMasterError(f"invalid {key} selection: {path.parent}")
        fields.append(tuple(values))
    if len(set(fields[-1])) != 50:
        raise ProtocolMasterError(f"duplicate valid-row ranks: {path.parent}")
    expected_offset = 25 if protocol == "o25" else 100
    if any(
        goal - start != expected_offset for start, goal in zip(fields[1], fields[2])
    ):
        raise ProtocolMasterError(f"selection offset mismatch: {path.parent}")
    return tuple(fields)


def _matrix(base: Path, *, protocol: str) -> tuple[dict[str, ScoreCell], list[Path]]:
    cells: dict[str, ScoreCell] = {}
    paths: list[Path] = []
    for key, relative in RESULT_PATHS.items():
        path = base / relative
        cells[key] = _load_result(path, protocol=protocol, expected_mode=key)
        paths.append(path)
    return cells, paths


def _row(label: str, loss: str, values: Mapping[str, ScoreCell]) -> MethodRow:
    return MethodRow(
        label=label,
        training_loss=loss,
        cells={key: values.get(key) for key in SCORE_KEYS},
    )


def _protocol_evidence(repository_root: Path, protocol: str) -> ProtocolMaster:
    artifacts = repository_root / "reports" / "artifacts"
    if protocol == "o25":
        old_root = artifacts / "actor_free_td_lewm_v1_c_c3_o25_20260906"
        c_base = (
            old_root / "v1_c_e10_o25_six_scores_3e36787_20260906" / "formal/o25/v1/c"
        )
        c3_path = (
            old_root
            / "v1_c3_e12_state_v_first_q2_a0p1_o25_d03a83b_20260906"
            / "formal/o25/v1/c3/state_v_plus_first_q2/alpha_0p1/results.json"
        )
        historical_runtime = "OSMesa*"
    elif protocol == "o100":
        old_root = artifacts / "actor_free_td_lewm_v1_c_c3_o100_20260906"
        c_base = old_root / "formal/o100/v1/c"
        c3_path = (
            old_root / "formal/o100/v1/c3/state_v_plus_first_q2/alpha_0p1/results.json"
        )
        historical_runtime = "OSMesa"
    else:
        raise ProtocolMasterError(f"unsupported protocol: {protocol}")

    c_cells, c_paths = _matrix(c_base, protocol=protocol)
    c3_cell = _load_result(
        c3_path,
        protocol=protocol,
        expected_mode="state_v_plus_first_q2",
    )
    c4_base = (
        artifacts
        / "actor_free_td_lewm_v1_c4_objective1_20260907"
        / f"formal/results/{protocol}/v1/c4"
    )
    c4_cells, c4_paths = _matrix(c4_base, protocol=protocol)

    all_paths = [*c_paths, c3_path, *c4_paths]
    selections = {_selection_signature(path, protocol=protocol) for path in all_paths}
    if len(selections) != 1:
        raise ProtocolMasterError(
            f"{protocol} results do not share one ordered selection"
        )
    selection = next(iter(selections))

    return ProtocolMaster(
        protocol=protocol,
        selection=selection[-1],
        rows=(
            _row(f"V1-C E10 / {historical_runtime}", "L_C", c_cells),
            _row(
                f"V1-C3 E12 / {historical_runtime}",
                "L_C3",
                {"state_v_plus_first_q2": c3_cell},
            ),
            _row("V1-C4 objective-v1 E10 / EGL", "L_C4=L_vector+L_goal", c4_cells),
        ),
    )


def load_master_evidence(repository_root: str | Path) -> MasterEvidence:
    root = Path(repository_root).expanduser().resolve()
    return MasterEvidence(
        o25=_protocol_evidence(root, "o25"),
        o100=_protocol_evidence(root, "o100"),
    )


def _rate(count: int) -> str:
    return f"{count}/50 ({count * 2}%)"


def _winner_sets(master: ProtocolMaster) -> tuple[list[int], list[set[int]], list[int]]:
    row_maxima: list[int] = []
    for row in master.rows:
        values = [cell.count for cell in row.cells.values() if cell is not None]
        if not values:
            raise ProtocolMasterError(f"empty method row: {row.label}")
        row_maxima.append(max(values))
    column_winners: list[set[int]] = []
    column_maxima: list[int] = []
    for key in SCORE_KEYS:
        available = [
            (index, row.cells[key].count)
            for index, row in enumerate(master.rows)
            if row.cells[key] is not None
        ]
        if not available:
            column_maxima.append(-1)
            column_winners.append(set())
            continue
        maximum = max(value for _, value in available)
        column_maxima.append(maximum)
        column_winners.append({index for index, value in available if value == maximum})
    return row_maxima, column_winners, column_maxima


def _markdown_table(master: ProtocolMaster) -> str:
    row_maxima, column_winners, _ = _winner_sets(master)
    headers = ["方法 / checkpoint / runtime", "训练 loss", *SCORE_HEADERS]
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row_index, row in enumerate(master.rows):
        values = [row.label, f"`{row.training_loss}`"]
        for column, key in enumerate(SCORE_KEYS):
            cell = row.cells[key]
            if cell is None:
                values.append("—")
                continue
            value = _rate(cell.count)
            row_best = cell.count == row_maxima[row_index]
            column_best = row_index in column_winners[column]
            if row_best:
                value = f"**{value}**"
            if column_best:
                value = "◆ " + value
            values.append(value)
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def _section(evidence: MasterEvidence) -> str:
    return "\n".join(
        [
            SECTION_START,
            "## O25 和 O100 当前正式结果独立总表",
            "",
            (
                "下面两张表各自使用固定的 50 个 start-goal pairs。每张表包含当前未被 "
                "supersede 的 13 个 C 系列正式单元：V1-C 六种评分、V1-C3 一个 "
                "State-V+First-Q2 评分、V1-C4 objective-v1 六种评分。C3 没有独立的 "
                "F-only 单元，因此写 `—`；同 EGL 的 V1-C F-only 后端审计重跑也不重复计数。"
            ),
            "",
            (
                "`◆` 表示该评分列的最高观察值，粗体表示该方法行的最高观察值。"
                "C4 objective-v0 已被新目标取代，只保留在历史表，不进入这里的赢家计算。"
            ),
            "",
            "### O25 当前正式结果总表",
            "",
            _markdown_table(evidence.o25),
            "",
            (
                "O25 的最高观察值为 38/50 (76%)：V1-C3 State-V+First-Q2 与 "
                "V1-C4 F+C4 tail 并列。C4 tail 在其同 EGL F-only 基线上 New=5、Lost=4，"
                "净增 1 个成功。"
            ),
            "",
            "### O100 当前正式结果总表",
            "",
            _markdown_table(evidence.o100),
            "",
            (
                "O100 的最高观察值是 V1-C First-Q 的 32/50 (64%)。V1-C4 的最高值为 "
                "Mean-Q 与 First-Q2 的 28/50 (56%)；二者相对同 EGL F-only 均净增 3 个成功。"
            ),
            "",
            (
                "运行口径：O100 的 V1-C/V1-C3 日志明确记录 OSMesa，C4 objective-v1 "
                "记录 EGL。O25 的 V1-C 历史运行由独立审计确认为 OSMesa 系列，但归档的 "
                "C3 O25 结果文件本身没有 renderer 字段，表中用 `OSMesa*` 标记这一限制。"
                "因此跨运行后端的横向差异只作描述；每个方法内部相对其 F-only 的配对结论才是主结论。"
            ),
            SECTION_END,
        ]
    )


def update_markdown_text(text: str, evidence: MasterEvidence) -> str:
    if SECTION_START in text or SECTION_END in text:
        raise ProtocolMasterError(
            "Markdown already contains current O25/O100 master tables"
        )
    anchor = "## 训练 / validation loss 证据"
    if text.count(anchor) != 1:
        raise ProtocolMasterError("Markdown training-loss anchor changed")
    opening = "模型均不训练 Actor。"
    replacement = (
        "模型均不训练 Actor。另有 O25 和 O100 两张独立当前正式总表，"
        "各覆盖 13 个未被 supersede 的 C 系列单元；它们不并入 O50 排名。"
    )
    if text.count(opening) != 1:
        raise ProtocolMasterError("Markdown opening paragraph changed")
    old_o25_heading = "## O25 配对补测 V1 C 与 C3"
    if text.count(old_o25_heading) != 1:
        raise ProtocolMasterError("Markdown O25 paired-detail heading changed")
    text = text.replace(opening, replacement, 1)
    text = text.replace(
        old_o25_heading,
        "## O25 V1 C 与 C3 配对明细",
        1,
    )
    return text.replace(anchor, _section(evidence) + "\n\n" + anchor, 1)


def _docx_helpers() -> tuple[Any, ...]:
    try:
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as error:
        raise RuntimeError(
            "python-docx is required from the bundled workspace runtime"
        ) from error
    return (
        WD_CELL_VERTICAL_ALIGNMENT,
        WD_ALIGN_PARAGRAPH,
        OxmlElement,
        qn,
        Inches,
        Pt,
        RGBColor,
    )


def _shade(cell: Any, fill: str) -> None:
    _, _, OxmlElement, qn, *_ = _docx_helpers()
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def _border(cell: Any, color: str = "D9D9D9", size: int = 4) -> None:
    _, _, OxmlElement, qn, *_ = _docx_helpers()
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _cell_text(
    cell: Any,
    text: str,
    *,
    bold: bool = False,
    color: str = "111827",
    size: float = 7.0,
) -> None:
    _, WD_ALIGN_PARAGRAPH, _, qn, _, Pt, RGBColor = _docx_helpers()
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Aptos")


def _new_docx_table(document: Any, master: ProtocolMaster) -> Any:
    (
        WD_CELL_VERTICAL_ALIGNMENT,
        _,
        OxmlElement,
        qn,
        Inches,
        _,
        _,
    ) = _docx_helpers()
    table = document.add_table(rows=1, cols=9)
    table.style = "Table Grid"
    table.autofit = False
    widths = (1.55, 1.15, 0.68, 0.76, 0.85, 0.75, 0.72, 0.75, 1.25)
    headers = (
        "Method / checkpoint / runtime",
        "Training loss",
        "F-only",
        "G/C4-only",
        "F+G/F+C4 tail",
        "First-Q\nalpha=.25",
        "Mean-Q",
        "First-Q2\nalpha=.25",
        "C3 State-V+First-Q2\nalpha=.10",
    )
    row_maxima, column_winners, _ = _winner_sets(master)
    for cell, value, width in zip(table.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        _cell_text(cell, value, bold=True, color="FFFFFF", size=6.8)
        _shade(cell, "17365D")
        _border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)

    for row_index, method in enumerate(master.rows):
        row = table.add_row()
        base_fill = "FFFFFF" if row_index % 2 == 0 else "F3F6FA"
        values: list[tuple[str, bool, bool]] = [
            (method.label, False, False),
            (method.training_loss, False, False),
        ]
        for column, key in enumerate(SCORE_KEYS):
            cell = method.cells[key]
            if cell is None:
                values.append(("—", False, False))
            else:
                values.append(
                    (
                        _rate(cell.count),
                        cell.count == row_maxima[row_index],
                        row_index in column_winners[column],
                    )
                )
        for cell, (value, row_best, column_best), width in zip(
            row.cells, values, widths
        ):
            cell.width = Inches(width)
            fill = base_fill
            if row_best and column_best:
                fill = "B7DEE8"
            elif column_best:
                fill = "DDEBF7"
            elif row_best:
                fill = "FFF2CC"
            _cell_text(cell, value, bold=row_best or column_best, size=6.8)
            _shade(cell, fill)
            _border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row_properties = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_properties.append(cant_split)
    return table


def _new_paragraph(
    document: Any, text: str, *, style: str = "Normal", bold: bool = False
) -> Any:
    _, _, _, qn, _, Pt, RGBColor = _docx_helpers()
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos Display" if style.startswith("Heading") else "Aptos"
    run.font.size = Pt(12.5 if style == "Heading 2" else 9.5)
    run.font.color.rgb = (
        RGBColor(0, 0, 0)
        if style.startswith("Heading")
        else RGBColor.from_string("111827")
    )
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), run.font.name
    )
    return paragraph


def _find_paragraph(document: Any, text: str) -> Any:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise ProtocolMasterError(
            f"DOCX paragraph {text!r} matched {len(matches)} times"
        )
    return matches[0]


def _replace_paragraph(document: Any, old: str, new: str) -> None:
    paragraph = _find_paragraph(document, old)
    nonempty = [run for run in paragraph.runs if run.text]
    if not nonempty:
        raise ProtocolMasterError(f"DOCX paragraph {old!r} has no text run")
    nonempty[0].text = new
    for run in nonempty[1:]:
        run.text = ""


def _insert_master_before(
    document: Any, anchor: Any, master: ProtocolMaster, marker: str
) -> None:
    heading = _new_paragraph(
        document,
        f"{master.protocol.upper()} current formal results master table",
        style="Heading 2",
    )
    body = _new_paragraph(
        document,
        (
            "Current non-superseded C-series evidence: 13 formal cells. C3 has no "
            "standalone F-only rerun; dashes are missing cells. Objective-v0 and "
            "same-EGL F-only audit reruns are excluded."
        ),
    )
    table = _new_docx_table(document, master)
    runtime_note = (
        "Runtime: historical O25 V1-C is from the OSMesa series and C4 "
        "objective-v1 = EGL. The archived O25 C3 result does not record its "
        "renderer."
        if master.protocol == "o25"
        else "Runtime: O100 V1-C/C3 logs = OSMesa and C4 objective-v1 = EGL."
    )
    note = _new_paragraph(
        document,
        (
            "Color key: yellow = row best; blue = score-column best; teal = both; "
            f"all ties are retained. {runtime_note} Cross-runtime comparisons are "
            "descriptive."
        ),
    )
    marker_paragraph = _new_paragraph(document, marker, bold=True)
    for element in (heading._p, body._p, table._tbl, note._p, marker_paragraph._p):
        anchor._p.addprevious(element)


def update_docx_document(document: Any, evidence: MasterEvidence) -> None:
    paragraph_text = {paragraph.text for paragraph in document.paragraphs}
    if DOCX_O25_MARKER in paragraph_text or DOCX_O100_MARKER in paragraph_text:
        raise ProtocolMasterError(
            "DOCX already contains current O25/O100 master tables"
        )
    if "RESULTS TD / V1-C4 FORMAL EXTENSION END" not in paragraph_text:
        raise ProtocolMasterError("DOCX is missing the completed C4 formal extension")
    _replace_paragraph(
        document,
        "Cube O25 paired audit for V1-C and V1-C3",
        "Cube O25 current formal results and paired details",
    )
    o25_anchor = _find_paragraph(
        document, "Methods by inference score, relative to F-only"
    )
    _replace_paragraph(
        document,
        "Methods by inference score, relative to F-only",
        "V1 C and C3 paired details relative to F only",
    )
    _insert_master_before(document, o25_anchor, evidence.o25, DOCX_O25_MARKER)

    _replace_paragraph(
        document,
        "Cube O100 paired evaluation: V1-C and V1-C3",
        "Cube O100 current formal results and paired details",
    )
    o100_anchor = _find_paragraph(document, "O100 method-by-inference-score matrix")
    _replace_paragraph(
        document,
        "O100 method-by-inference-score matrix",
        "V1 C and C3 paired details relative to F only",
    )
    _insert_master_before(document, o100_anchor, evidence.o100, DOCX_O100_MARKER)


def _atomic_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    temporary = Path(name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def write_updated_reports(
    *,
    docx_path: str | Path,
    markdown_path: str | Path,
    repository_root: str | Path,
    output_dir: str | Path,
) -> dict[str, str]:
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError(
            "python-docx is required from the bundled workspace runtime"
        ) from error
    source_docx = Path(docx_path).expanduser().resolve()
    source_markdown = Path(markdown_path).expanduser().resolve()
    destination = Path(output_dir).expanduser().resolve()
    destination.mkdir(parents=True, exist_ok=True)
    evidence = load_master_evidence(repository_root)
    markdown = update_markdown_text(
        source_markdown.read_text(encoding="utf-8"), evidence
    )
    document = Document(source_docx)
    update_docx_document(document, evidence)
    output_docx = destination / source_docx.name
    output_markdown = destination / source_markdown.name
    document.save(output_docx)
    _atomic_text(output_markdown, markdown)
    return {"docx": str(output_docx), "markdown": str(output_markdown)}


__all__ = [
    "DOCX_O100_MARKER",
    "DOCX_O25_MARKER",
    "MasterEvidence",
    "ProtocolMasterError",
    "SECTION_END",
    "SECTION_START",
    "SCORE_HEADERS",
    "SCORE_KEYS",
    "load_master_evidence",
    "update_docx_document",
    "update_markdown_text",
    "write_updated_reports",
]
