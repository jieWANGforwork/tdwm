#!/usr/bin/env python3
"""Append the validated V1-C/V1-C3 O100 paired matrix to Results TD.

The script updates the existing canonical report in place.  It deliberately
keeps O100 separate from the O25/O50 tables because the goal offset and
execution budget differ, while preserving the same method-by-score layout.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import io
import json
import os
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from types import ModuleType
from typing import Any, Mapping, Sequence

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = REPOSITORY_ROOT.parents[1]
DEFAULT_ARTIFACT_ROOT = (
    PROJECT_ROOT / "artifacts/v1_c_c3_o100_b77aabd_20260906_attempt02"
)
DEFAULT_REPO_DOCX = (
    REPOSITORY_ROOT
    / "reports/results_td_actor_free_td_lewm_complete_cube_seed3072.docx"
)
DEFAULT_PROJECT_DOCX = PROJECT_ROOT / "Results TD.docx"

DOCX_START = "RESULTS TD / O100 V1-C C3 PAIRED EXTENSION"
DOCX_END = "RESULTS TD / O100 V1-C C3 PAIRED EXTENSION END"
EPISODES = 50
SELECTION_SHA256 = "8a87815e8e1816ccb5021af81a5e2307a5b342d094eec3edf221a0e24851d10c"
ACTION_SHA256 = "57f4d3c252e1805f4af1f614d20d1d1a064fa0d1d463ed5eb8ecf9dfc2b1a723"
V1_C_CHECKPOINT_SHA256 = "88bd65c48a6c701852f50552ec8f9109d6ae8ac57c467de207aa2c652c0f59a3"
V1_C3_CHECKPOINT_SHA256 = "5e240053d7c33fc016ef2ff64f3a4a79706dbe10dfde347d5c5f3cd45043e5b2"


@dataclass(frozen=True)
class CellSpec:
    key: str
    relative: Path
    score_mode: str
    variant: str
    alpha: float | None
    epoch: int
    success_count: int
    results_sha256: str


CELL_SPECS = (
    CellSpec(
        "f_only",
        Path("formal/o100/v1/c/f_only"),
        "f_only",
        "c",
        None,
        10,
        25,
        "498248dd34cf94c42b09bbf25783aa322cc5058babab3e79801ae6c943ec549f",
    ),
    CellSpec(
        "g_only",
        Path("formal/o100/v1/c/g_only"),
        "g_only",
        "c",
        None,
        10,
        24,
        "c1a8fedf2d939c242cec7ec9cba1e44753729edcceb5c7ea7636b123f58a6e58",
    ),
    CellSpec(
        "f_plus_g",
        Path("formal/o100/v1/c/f_plus_g"),
        "f_plus_g",
        "c",
        None,
        10,
        22,
        "ab29462a06af08e462d364551646a642934024d115d7063df6e23c19e6225e2c",
    ),
    CellSpec(
        "first_q",
        Path("formal/o100/v1/c/f_plus_g_first/alpha_0p25"),
        "f_plus_g_first",
        "c",
        0.25,
        10,
        32,
        "f7a4fb4d59b676cf9e8b878f4afcfa7fbdc29949a0d7c12168a169e61099513a",
    ),
    CellSpec(
        "mean_q",
        Path("formal/o100/v1/c/g_only_f_rollout_mean"),
        "g_only_f_rollout_mean",
        "c",
        None,
        10,
        25,
        "a0e697dae65f99706ba64ffad6cd572f14e314c6cda557fb5dfbd3df3f01ddf4",
    ),
    CellSpec(
        "first_q2",
        Path("formal/o100/v1/c/f_plus_g_first_q2/alpha_0p25"),
        "f_plus_g_first_q2",
        "c",
        0.25,
        10,
        26,
        "decbaf864e5ccb973f71fe54aeaf3c7d99e33f5f21d6987eca8f6426a46a7194",
    ),
    CellSpec(
        "c3",
        Path("formal/o100/v1/c3/state_v_plus_first_q2/alpha_0p1"),
        "state_v_plus_first_q2",
        "c3",
        0.10,
        12,
        25,
        "e185267f61ae349116d86903c537b4ddcae07b55b23a8b389087266fe9a26d5a",
    ),
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_json(path: Path) -> Mapping[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, Mapping):
        raise ValueError(f"{path} must contain a JSON object")
    return value


def _require(actual: Any, expected: Any, context: str) -> None:
    if actual != expected:
        raise ValueError(f"{context}: expected {expected!r}, found {actual!r}")


def _load_outcomes(artifact_root: Path) -> dict[str, tuple[bool, ...]]:
    outcomes: dict[str, tuple[bool, ...]] = {}
    shared_selection: Mapping[str, Any] | None = None
    shared_action: Mapping[str, Any] | None = None
    for spec in CELL_SPECS:
        directory = artifact_root / spec.relative
        results_path = directory / "results.json"
        protocol_path = directory / "protocol_manifest.json"
        selection_path = directory / "episode_selection.json"
        action_path = directory / "action_normalization.json"
        for path in (results_path, protocol_path, selection_path, action_path):
            if not path.is_file():
                raise FileNotFoundError(path)
        _require(_sha256(results_path), spec.results_sha256, f"{spec.key} results SHA-256")
        _require(_sha256(selection_path), SELECTION_SHA256, f"{spec.key} selection SHA-256")
        _require(_sha256(action_path), ACTION_SHA256, f"{spec.key} action SHA-256")

        result = _read_json(results_path)
        protocol = _read_json(protocol_path)
        selection = _read_json(selection_path)
        action = _read_json(action_path)
        if shared_selection is None:
            shared_selection = selection
            shared_action = action
        else:
            _require(selection, shared_selection, f"{spec.key} shared selection")
            _require(action, shared_action, f"{spec.key} shared action normalization")

        for values, label in ((result, "results"), (protocol, "protocol")):
            for key, expected in {
                "evaluation_protocol": "O100",
                "protocol_label": "o100",
                "goal_offset": 100,
                "episode_budget": 200,
                "score_mode": spec.score_mode,
                "g_first_weight": spec.alpha,
            }.items():
                _require(values.get(key), expected, f"{spec.key}.{label}.{key}")
        _require(result.get("variant"), spec.variant, f"{spec.key}.variant")
        checkpoint = protocol.get("checkpoint")
        if not isinstance(checkpoint, Mapping):
            raise ValueError(f"{spec.key}.checkpoint is missing")
        _require(checkpoint.get("epoch"), spec.epoch, f"{spec.key}.checkpoint.epoch")
        _require(
            checkpoint.get("sha256"),
            V1_C3_CHECKPOINT_SHA256 if spec.variant == "c3" else V1_C_CHECKPOINT_SHA256,
            f"{spec.key}.checkpoint.sha256",
        )
        metrics = result.get("metrics")
        if not isinstance(metrics, Mapping):
            raise ValueError(f"{spec.key}.metrics is missing")
        flags = metrics.get("episode_successes")
        if (
            not isinstance(flags, list)
            or len(flags) != EPISODES
            or any(type(flag) is not bool for flag in flags)
        ):
            raise ValueError(f"{spec.key} must contain exactly 50 Boolean outcomes")
        _require(sum(flags), spec.success_count, f"{spec.key} success count")
        outcomes[spec.key] = tuple(flags)
    return outcomes


def _pair_ids(indices: set[int]) -> str:
    return ", ".join(f"P{index + 1:02d}" for index in sorted(indices))


def _stats(
    baseline: Sequence[bool], candidate: Sequence[bool]
) -> tuple[int, set[int], set[int], int]:
    new = {
        index
        for index, (old, current) in enumerate(zip(baseline, candidate))
        if not old and current
    }
    lost = {
        index
        for index, (old, current) in enumerate(zip(baseline, candidate))
        if old and not current
    }
    return sum(candidate), new, lost, sum(baseline) + len(new)


def _score(count: int) -> str:
    return f"{count}/50 ({2 * count}%)"


def _delta(new: set[int], lost: set[int], f_plus_new: int) -> str:
    return f"{len(new)} / {len(lost)} / {f_plus_new}"


def _load_helpers() -> ModuleType:
    path = REPOSITORY_ROOT / "scripts/build_results_td_v1.py"
    spec = importlib.util.spec_from_file_location("_results_td_o100_helpers", path)
    if spec is None or spec.loader is None:
        raise ImportError(path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _set_table_borders(table: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        for cell in row.cells:
            properties = cell._tc.get_or_add_tcPr()
            borders = properties.find(qn("w:tcBorders"))
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                properties.append(borders)
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                element = borders.find(qn(f"w:{edge}"))
                if element is None:
                    element = OxmlElement(f"w:{edge}")
                    borders.append(element)
                element.set(qn("w:val"), "single")
                element.set(qn("w:sz"), "4")
                element.set(qn("w:color"), "D9D9D9")


def _format_table(table: Any, helpers: ModuleType) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    for cell in table.rows[0].cells:
        helpers._shade_cell(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string("000000")
            run.font.size = Pt(7.0)
            run.bold = True
    for row in table.rows[1:]:
        for index, cell in enumerate(row.cells):
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            if index > 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(7.2)
    _set_table_borders(table)


def _remove_existing(document: Any) -> bool:
    starts = [p for p in document.paragraphs if p.text == DOCX_START]
    ends = [p for p in document.paragraphs if p.text == DOCX_END]
    if not starts and not ends:
        return False
    if len(starts) != 1 or len(ends) != 1:
        raise ValueError("O100 DOCX markers are incomplete or duplicated")
    body = document._element.body
    children = list(body)
    start_index = children.index(starts[0]._p)
    end_index = children.index(ends[0]._p)
    if end_index < start_index:
        raise ValueError("O100 DOCX markers are reversed")
    for child in children[start_index : end_index + 1]:
        body.remove(child)
    return True


def _set_running_matter(section: Any, helpers: ModuleType) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for header in (section.header, section.even_page_header, section.first_page_header):
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        helpers._set_run_font(
            paragraph.add_run("Results TD · Cube O100 · V1-C and V1-C3 paired evaluation"),
            size=8.5,
            color="6B7280",
        )
    for footer in (section.footer, section.even_page_footer, section.first_page_footer):
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        helpers._set_run_font(
            paragraph.add_run("Validated paired outcomes · Page "),
            size=8.5,
            color="6B7280",
        )
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        paragraph._p.append(field)


def _build_payload(
    source_docx: Path, outcomes: Mapping[str, Sequence[bool]]
) -> bytes:
    from docx import Document
    from docx.shared import Pt

    helpers = _load_helpers()
    document = Document(str(source_docx))
    existed = _remove_existing(document)
    if not existed:
        helpers._configure_append_section(document)
    _set_running_matter(document.sections[-1], helpers)

    baseline = outcomes["f_only"]
    cell_stats = {
        key: _stats(baseline, flags) for key, flags in outcomes.items()
    }
    _, first_new, _, _ = cell_stats["first_q"]
    other_keys = ("g_only", "f_plus_g", "mean_q", "first_q2", "c3")
    other_new = set().union(*(cell_stats[key][1] for key in other_keys))
    all_new = first_new | other_new

    marker = document.add_paragraph(style="Report Kicker")
    helpers._set_run_font(
        marker.add_run(DOCX_START), size=9.5, color="5C6975", bold=True
    )
    title = document.add_heading("Cube O100 paired evaluation: V1-C and V1-C3", level=1)
    title.paragraph_format.page_break_before = False
    for run in title.runs:
        helpers._set_run_font(run, size=22, color="000000", bold=True)
    subtitle = document.add_paragraph()
    helpers._set_run_font(
        subtitle.add_run(
            "One fixed set of 50 start-goal pairs: score, New, Lost and complementary successes"
        ),
        size=12,
        color="4B5563",
    )
    intro = document.add_paragraph(
        "All columns use the same O100 pairs, planning seed 42, CEM 300 x 30 and budget 200. "
        "V1-C uses E10 and "
        "V1-C3 uses E12. The best standalone score is First-Q alpha=.25 at 32/50 (64%): "
        "seven New and zero Lost relative to F-only at 25/50 (50%)."
    )
    for run in intro.runs:
        helpers._set_run_font(run, size=9.5, color="7A5A00", bold=True)

    document.add_heading("O100 method-by-inference-score matrix", level=2)
    document.add_paragraph(
        "Every score is followed by New / Lost / F+New counts. New means the scorer succeeds "
        "where F-only fails; Lost means an F-only success becomes a failure; F+New is the "
        "post-hoc upper bound that keeps every F success and adds all New. The C3 F-only cell "
        "is the same frozen-F baseline, not another evaluation run."
    )

    def values(key: str) -> tuple[str, str]:
        count, new, lost, f_plus_new = cell_stats[key]
        return _score(count), _delta(new, lost, f_plus_new)

    g_score, g_change = values("g_only")
    tail_score, tail_change = values("f_plus_g")
    first_score, first_change = values("first_q")
    mean_score, mean_change = values("mean_q")
    first2_score, first2_change = values("first_q2")
    c3_score, c3_change = values("c3")
    rows = (
        (
            "V1-C E10",
            _score(sum(baseline)),
            g_score,
            g_change,
            tail_score,
            tail_change,
            first_score,
            first_change,
            mean_score,
            mean_change,
            first2_score,
            first2_change,
            "—",
            "—",
        ),
        (
            "V1-C3 E12",
            _score(sum(baseline)) + "*",
            "—",
            "—",
            "—",
            "—",
            "—",
            "—",
            "—",
            "—",
            "—",
            "—",
            c3_score,
            c3_change,
        ),
    )
    table = helpers._add_table(
        document,
        headers=(
            "Method / checkpoint",
            "F-only\nbaseline",
            "G-only",
            "New / Lost / F+New",
            "F+G tail",
            "New / Lost / F+New",
            "First-Q\nalpha=.25",
            "New / Lost / F+New",
            "Mean-Q",
            "New / Lost / F+New",
            "First-Q2\nalpha=.25",
            "New / Lost / F+New",
            "C3 State-V + First-Q2\nalpha=.10",
            "New / Lost / F+New",
        ),
        rows=rows,
        widths=(1800, 1100, 950, 800, 950, 800, 1050, 800, 950, 800, 1050, 800, 1600, 950),
    )
    _format_table(table, helpers)
    helpers._shade_cell(table.rows[1].cells[6], "B7DEE8")
    helpers._shade_cell(table.rows[1].cells[7], "B7DEE8")
    legend = document.add_paragraph(
        "Cyan marks the best standalone O100 result; it also preserves every F-only success."
    )
    for run in legend.runs:
        helpers._set_run_font(run, size=8.3, color="5C6975")

    document.add_heading("First-Q's seven New pairs and complementary rescues", level=2)
    conclusions = (
        f"First-Q has seven New pairs: {_pair_ids(first_new)}. Lost=0, so its observed 32/50 equals its F+New upper bound.",
        "Other scores do rescue four distinct pairs outside those seven: P13 only by F+G tail; "
        "P25 only by G-only; P42 only by First-Q2; and P48 by both F+G tail and Mean-Q.",
        f"The post-hoc union of F-only, First-Q and every complementary rescue is "
        f"{sum(baseline) + len(all_new)}/50 ({2 * (sum(baseline) + len(all_new))}%). "
        "This is an oracle selected with true success labels, not a deployable result.",
        "C3's three New pairs (P29, P31 and P40) are already included in First-Q's seven; "
        "C3 contributes no rescue outside First-Q.",
        "The next target is a conservative gate around First-Q: retain First-Q by default, and "
        "switch to tail, G-only or First-Q2 only when a separate development set can reliably "
        "identify P13-, P25-, P42- or P48-like cases. Lock the rule before testing unseen pairs "
        "and multiple planning seeds.",
        "Scope: one training seed, one planning seed and one fixed set of 50 formal O100 pairs. "
        "The 36/50 (72%) oracle is not an implemented controller result.",
    )
    for text in conclusions:
        paragraph = document.add_paragraph(text)
        paragraph.style = document.styles["List Bullet"]

    evidence = document.add_paragraph(
        "Evidence: seven scores x 50 pair-level outcomes; shared selection SHA-256 "
        f"{SELECTION_SHA256}。"
    )
    for run in evidence.runs:
        helpers._set_run_font(run, size=8.3, color="5C6975")

    end = document.add_paragraph()
    end_run = end.add_run(DOCX_END)
    end_run.font.hidden = True
    end_run.font.size = Pt(1)

    stream = io.BytesIO()
    document.save(stream)
    payload = stream.getvalue()
    if not payload.startswith(b"PK"):
        raise RuntimeError("python-docx did not produce OOXML")
    return payload


def _atomic_write(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="wb",
        dir=path.parent,
        prefix=f".{path.name}.",
        suffix=".tmp",
        delete=False,
    ) as stream:
        stream.write(payload)
        stream.flush()
        os.fsync(stream.fileno())
        temporary = Path(stream.name)
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Append the V1-C/V1-C3 O100 paired matrix to the existing Results TD"
    )
    parser.add_argument("--artifact-root", type=Path, default=DEFAULT_ARTIFACT_ROOT)
    parser.add_argument("--repo-docx", type=Path, default=DEFAULT_REPO_DOCX)
    parser.add_argument("--project-docx", type=Path, default=DEFAULT_PROJECT_DOCX)
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args(argv)

    artifact_root = args.artifact_root.expanduser().resolve()
    repo_docx = args.repo_docx.expanduser().resolve()
    project_docx = args.project_docx.expanduser().resolve()
    outcomes = _load_outcomes(artifact_root)
    baseline = outcomes["f_only"]
    first = _stats(baseline, outcomes["first_q"])
    other_new = set().union(
        *(
            _stats(baseline, outcomes[key])[1]
            for key in ("g_only", "f_plus_g", "mean_q", "first_q2", "c3")
        )
    )
    outside = other_new - first[1]
    if args.validate_only:
        print(
            "PASS: O100 7/7 cells; F-only=25/50; First-Q=32/50; "
            f"outside-First-Q rescues={_pair_ids(outside)}; oracle=36/50"
        )
        return 0

    if not project_docx.is_file():
        raise FileNotFoundError("the canonical project Results TD DOCX is required")
    # The project-root document is the user-facing canonical copy and can be
    # newer than the tracked report mirror.  Always preserve it as the source;
    # after the append, synchronize the resulting payload back to the mirror.
    payload = _build_payload(project_docx, outcomes)
    _atomic_write(repo_docx, payload)
    _atomic_write(project_docx, payload)
    print(
        "PASS: updated the existing Results TD DOCX in place; "
        "F-only=25/50; First-Q=32/50; F|First-Q|others=36/50"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
