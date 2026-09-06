#!/usr/bin/env python3
"""Validate and append the paired V1-C/V1-C3 Cube O25 result extension.

This is intentionally downstream of the formal evaluators.  It reads only the
archived JSON evidence, fails closed on every result/protocol fingerprint, and
never invokes the historical Results TD builders.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import math
import os
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from types import ModuleType
from typing import Any, Mapping, Sequence

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
ARTIFACT_ROOT = (
    REPOSITORY_ROOT
    / "reports/artifacts/actor_free_td_lewm_v1_c_c3_o25_20260906"
)
MARKDOWN_PATH = (
    REPOSITORY_ROOT / "reports/actor_free_td_lewm_complete_cube_seed3072.md"
)
DOCX_PATH = (
    REPOSITORY_ROOT
    / "reports/results_td_actor_free_td_lewm_complete_cube_seed3072.docx"
)
PROJECT_DOCX_PATH = REPOSITORY_ROOT.parents[1] / "Results TD.docx"
PAIRED_PATH = ARTIFACT_ROOT / "paired_outcomes.csv"
RECONCILIATION_PATH = ARTIFACT_ROOT / "reconciliation_ledger.json"

MARKDOWN_START = "<!-- RESULTS_TD_O25_V1_C_C3_START -->"
MARKDOWN_END = "<!-- RESULTS_TD_O25_V1_C_C3_END -->"
DOCX_START = "RESULTS TD / O25 PAIRED EXTENSION"
DOCX_END = "RESULTS TD / O25 PAIRED EXTENSION END"

EPISODES = 50
SELECTION_SHA256 = "56546fe8725ce0e4670f308c5b325bd64ff2a792373add8c20ddbcab02da6b37"
RANKS_SHA256 = "72af45d4bad65a25288c5d405072d18ab5c0b4f0b67ddc970ac3f344b3c22fd9"
ACTION_SHA256 = "57f4d3c252e1805f4af1f614d20d1d1a064fa0d1d463ed5eb8ecf9dfc2b1a723"
V1_C_CHECKPOINT_SHA256 = "88bd65c48a6c701852f50552ec8f9109d6ae8ac57c467de207aa2c652c0f59a3"
V1_C3_CHECKPOINT_SHA256 = "5e240053d7c33fc016ef2ff64f3a4a79706dbe10dfde347d5c5f3cd45043e5b2"
BASE_MARKDOWN_SHA256 = "46956daf27fe7ab6db4746176a2d64079ee4d463f824ad387fdef940a3b3f14b"
BASE_DOCX_SHA256 = "a6c9c30d9ace1e109b06bdee46fb5d58c11bf01343729cc80b47c90d28c8e86a"

V1_RUN = Path("v1_c_e10_o25_six_scores_3e36787_20260906")
C3_RUN = Path("v1_c3_e12_state_v_first_q2_a0p1_o25_d03a83b_20260906")


@dataclass(frozen=True)
class CellSpec:
    key: str
    label: str
    directory: Path
    results_sha256: str
    protocol_sha256: str
    method: str
    variant: str
    score_mode: str
    alpha: float | None
    checkpoint_sha256: str
    epoch: int
    global_step: int
    horizon: int
    receding_horizon: int
    success_count: int
    historical_o50_count: int


CELL_SPECS = (
    CellSpec("f_only", "F-only", V1_RUN / "formal/o25/v1/c/f_only", "81bf85ed3f894be788edece44fe036eca10229d0943030f558f807137c28ef51", "e46dddbbb5f23c9aa4c815623c0553f944ba10a05374318127c89d0d3b0770a7", "actor_free_td_lewm_v1_c", "c", "f_only", None, V1_C_CHECKPOINT_SHA256, 10, 127960, 5, 5, 37, 23),
    CellSpec("g_only", "G-only", V1_RUN / "formal/o25/v1/c/g_only", "8ff0265595f1903d5f522f16806bd45200357708ed7bf8b6b4a620ce586955c8", "444be50326f7885d6929335d0f11579fffc262b253a60a77f03348c86cab9ca9", "actor_free_td_lewm_v1_c", "c", "g_only", None, V1_C_CHECKPOINT_SHA256, 10, 127960, 1, 1, 29, 18),
    CellSpec("f_plus_g", "F+G tail", V1_RUN / "formal/o25/v1/c/f_plus_g", "9d48ffcb9da09946e7aa466b66d4d0b582053d3a4cdf5c7930959b2afaf36f2b", "28177f46bc60de6e05bbe8179a8f585bee09faf023b69af27217199ecb759d1a", "actor_free_td_lewm_v1_c", "c", "f_plus_g", None, V1_C_CHECKPOINT_SHA256, 10, 127960, 5, 5, 35, 22),
    CellSpec("first_q", "First-Q alpha=.25", V1_RUN / "formal/o25/v1/c/f_plus_g_first/alpha_0p25", "409ccb1668b8463f457668c667d9535fc183713bd5252c477f0bb4a33f771704", "112d1caa00fe17b9cc85ba89d29eabfcbb5ca0be0fcef784d6d7be19b757c4ad", "actor_free_td_lewm_v1_c", "c", "f_plus_g_first", 0.25, V1_C_CHECKPOINT_SHA256, 10, 127960, 5, 5, 36, 28),
    CellSpec("mean_q", "Mean-Q rollout", V1_RUN / "formal/o25/v1/c/g_only_f_rollout_mean", "a227717097f05e19fa697c7eaa5d0bb910a7fb050849ad7420c7a4fb80877e57", "c0d90cea44458486f2defa601e3018684dbced97a01f233be8d06aadfcec34e0", "actor_free_td_lewm_v1_c", "c", "g_only_f_rollout_mean", None, V1_C_CHECKPOINT_SHA256, 10, 127960, 5, 5, 30, 21),
    CellSpec("first_q2", "First-Q2 alpha=.25", V1_RUN / "formal/o25/v1/c/f_plus_g_first_q2/alpha_0p25", "5de7636df1e37a974578a0a453749aa5a91507dce247d83f61108fdc2cf31970", "0732a5f2ae69a0cac89023022ffdada8952796b4f965d608c484307606be68e4", "actor_free_td_lewm_v1_c", "c", "f_plus_g_first_q2", 0.25, V1_C_CHECKPOINT_SHA256, 10, 127960, 5, 5, 35, 26),
    CellSpec("c3", "C3 State-V + First-Q2 alpha=.10", C3_RUN / "formal/o25/v1/c3/state_v_plus_first_q2/alpha_0p1", "1bc87b98d31721d2cac59488723dfd96df8156858dd78126c6bb1df69f112b87", "a187482aeea1bd4f41a02ae5d598763cd57574f8e003b0033dd76e726b025356", "actor_free_td_lewm_v1_c3", "c3", "state_v_plus_first_q2", 0.1, V1_C3_CHECKPOINT_SHA256, 12, 12000, 5, 5, 38, 31),
)

RAW_FILE_HASHES = {
    V1_RUN / "checkpoint_manifest.json": "52ad4fc6bb137f8941d576e4dc4698757b30748334cd7fadfb3a74606db7d1bf",
    V1_RUN / "formal/_launcher/launcher_manifest.json": "9dd8e791a88a97885ad917a5b72a7c3826d423d9599661dd10d485e59ce1b650",
    C3_RUN / "checkpoint_manifest.json": "18b7c605d54320886772134c264ed118c2f63c167bc8eb178607e641ea993d45",
    C3_RUN / "formal/_launcher/launcher_manifest.json": "42b92d72dd251f28f56de711c1b6685c43780be20a071006682d035e043648db",
}

SCORE_PATHS = {
    "F-only": "J_F = ||z5^F - z_goal||^2",
    "G-only": "J_G = -Q_G(z0,A1,g); H=1",
    "F+G tail": "J_tail = ||z4^F-z_goal||^2 - gamma^4 Q_G(z4^F,A5,g)",
    "First-Q alpha=.25": "J = J_F - .25 Q_G(z0,A1,g)",
    "Mean-Q rollout": "J = -mean[k=1..5] Q_G(z{k-1}^F,Ak,g)",
    "First-Q2 alpha=.25": "J = Zcand(J_F) - .25 Zcand(Q_G(z0,A1,g))",
    "C3 State-V + First-Q2 alpha=.10": "J = Zcand(Vbar(F^5,z_goal)) - .10 Zcand(Q_G(z0,A1,g))",
}


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def file_sha256(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def canonical_json_sha256(value: Any) -> str:
    return sha256_bytes(json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode("utf-8"))


def read_json(path: Path) -> Mapping[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, Mapping):
        raise ValueError(f"{path} must contain a JSON object")
    return value


def require_equal(actual: Any, expected: Any, context: str) -> None:
    if actual != expected:
        raise ValueError(f"{context}: expected {expected!r}, found {actual!r}")


def pair_ids(flags: Sequence[bool]) -> list[str]:
    return [f"P{index + 1:02d}" for index, flag in enumerate(flags) if flag]


def _validate_launcher(path: Path, *, expected_jobs: int) -> None:
    launcher = read_json(path)
    for key, expected in {
        "status": "SUCCEEDED",
        "stage": "formal",
        "evaluation_protocol": "O25",
        "protocol_label": "o25",
        "training_performed": False,
        "inference_only": True,
        "alpha_selection_performed": False,
    }.items():
        require_equal(launcher.get(key), expected, f"{path}.{key}")
    jobs = launcher.get("jobs")
    if not isinstance(jobs, Mapping) or len(jobs) != expected_jobs:
        raise ValueError(f"{path}.jobs must contain exactly {expected_jobs} jobs")
    for job_id, job in jobs.items():
        if not isinstance(job, Mapping):
            raise ValueError(f"{path}.jobs.{job_id} must be an object")
        require_equal(job.get("state"), "SUCCEEDED", f"{path}.jobs.{job_id}.state")
        require_equal(job.get("exit_code"), 0, f"{path}.jobs.{job_id}.exit_code")
        evidence = job.get("evidence")
        if not isinstance(evidence, Mapping):
            raise ValueError(f"{path}.jobs.{job_id}.evidence is missing")
        require_equal(evidence.get("selection_file_sha256"), SELECTION_SHA256, f"{job_id}.selection")
        require_equal(evidence.get("valid_row_ranks_sha256"), RANKS_SHA256, f"{job_id}.ranks")
        require_equal(evidence.get("action_normalization_sha256"), ACTION_SHA256, f"{job_id}.action")


def validate_evidence(artifact_root: Path) -> tuple[dict[str, Any], Mapping[str, Any]]:
    for relative, expected_hash in RAW_FILE_HASHES.items():
        path = artifact_root / relative
        if not path.is_file():
            raise FileNotFoundError(path)
        require_equal(file_sha256(path), expected_hash, f"SHA-256 {relative}")
    _validate_launcher(artifact_root / V1_RUN / "formal/_launcher/launcher_manifest.json", expected_jobs=6)
    _validate_launcher(artifact_root / C3_RUN / "formal/_launcher/launcher_manifest.json", expected_jobs=1)

    outcomes: dict[str, Any] = {}
    shared_selection: Mapping[str, Any] | None = None
    shared_action: Mapping[str, Any] | None = None
    seen_results_hashes: set[str] = set()
    for spec in CELL_SPECS:
        directory = artifact_root / spec.directory
        paths = {name: directory / name for name in ("results.json", "protocol_manifest.json", "episode_selection.json", "action_normalization.json")}
        missing = [str(path) for path in paths.values() if not path.is_file()]
        if missing:
            raise FileNotFoundError(", ".join(missing))
        require_equal(file_sha256(paths["results.json"]), spec.results_sha256, f"{spec.key} results SHA-256")
        require_equal(file_sha256(paths["protocol_manifest.json"]), spec.protocol_sha256, f"{spec.key} protocol SHA-256")
        require_equal(file_sha256(paths["episode_selection.json"]), SELECTION_SHA256, f"{spec.key} selection SHA-256")
        require_equal(file_sha256(paths["action_normalization.json"]), ACTION_SHA256, f"{spec.key} action SHA-256")
        if spec.results_sha256 in seen_results_hashes:
            raise ValueError(f"duplicate result payload: {spec.key}")
        seen_results_hashes.add(spec.results_sha256)

        selection = read_json(paths["episode_selection.json"])
        action = read_json(paths["action_normalization.json"])
        if shared_selection is None:
            shared_selection = selection
            shared_action = action
        else:
            require_equal(selection, shared_selection, f"{spec.key} shared selection")
            require_equal(action, shared_action, f"{spec.key} shared action normalization")
        for name in ("episode_indices", "start_steps", "goal_steps", "valid_row_ranks"):
            values = selection.get(name)
            if not isinstance(values, list) or len(values) != EPISODES or any(type(value) is not int for value in values):
                raise ValueError(f"{spec.key}.{name} must contain 50 integers")
        if any(goal - start != 25 for start, goal in zip(selection["start_steps"], selection["goal_steps"])):
            raise ValueError(f"{spec.key} is not an exact O25 selection")
        require_equal(canonical_json_sha256(selection["valid_row_ranks"]), RANKS_SHA256, f"{spec.key} rank digest")

        result = read_json(paths["results.json"])
        protocol = read_json(paths["protocol_manifest.json"])
        for values, label in ((result, "results"), (protocol, "protocol")):
            for key, expected in {
                "evaluation_protocol": "O25", "protocol_label": "o25", "goal_offset": 25,
                "episode_budget": 50, "score_mode": spec.score_mode,
                "receding_horizon": spec.receding_horizon,
                "executed_action_blocks_before_replanning": spec.receding_horizon,
                "executed_environment_steps_before_replanning": spec.receding_horizon * 5,
                "executed_action_block": "first_block_only" if spec.receding_horizon == 1 else "all_five_blocks",
                "replanning": "every_action_block" if spec.receding_horizon == 1 else "every_five_action_blocks",
                "cem_execution": "execute_A1_from_minimum_total_cost_plan" if spec.receding_horizon == 1 else "execute_A1_through_A5_from_minimum_total_cost_plan",
            }.items():
                require_equal(values.get(key), expected, f"{spec.key}.{label}.{key}")
            require_equal(values.get("g_first_weight"), spec.alpha, f"{spec.key}.{label}.alpha")
        for key, expected in {
            "method": spec.method, "variant": spec.variant, "score_mode": spec.score_mode,
            "planning_horizon": spec.horizon, "smoke": False, "pilot": False,
        }.items():
            require_equal(result.get(key), expected, f"{spec.key}.results.{key}")
        metrics = result.get("metrics")
        if not isinstance(metrics, Mapping):
            raise ValueError(f"{spec.key}.metrics must be an object")
        flags = metrics.get("episode_successes")
        if not isinstance(flags, list) or len(flags) != EPISODES or any(type(flag) is not bool for flag in flags):
            raise ValueError(f"{spec.key} must contain exactly 50 Boolean outcomes")
        require_equal(sum(flags), spec.success_count, f"{spec.key} success count")
        if not math.isclose(float(metrics.get("success_rate")), spec.success_count * 2.0, abs_tol=1e-9):
            raise ValueError(f"{spec.key} success rate disagrees with its outcomes")
        checkpoint = protocol.get("checkpoint")
        if not isinstance(checkpoint, Mapping):
            raise ValueError(f"{spec.key}.checkpoint is missing")
        for key, expected in {"sha256": spec.checkpoint_sha256, "epoch": spec.epoch, "global_step": spec.global_step, "method": spec.method, "variant": spec.variant}.items():
            require_equal(checkpoint.get(key), expected, f"{spec.key}.checkpoint.{key}")
        require_equal(protocol.get("selection"), selection, f"{spec.key}.protocol.selection")
        normalization = protocol.get("normalization")
        if not isinstance(normalization, Mapping):
            raise ValueError(f"{spec.key}.normalization is missing")
        require_equal(normalization.get("action"), action, f"{spec.key}.protocol.action")
        configured = protocol.get("protocol")
        if not isinstance(configured, Mapping) or not isinstance(configured.get("planning"), Mapping):
            raise ValueError(f"{spec.key}.protocol.planning is missing")
        planning = configured["planning"]
        for key, expected in {"horizon": spec.horizon, "receding_horizon": spec.receding_horizon, "candidates": 300, "iterations": 30, "elites": 30, "action_block": 5, "planning_seed": 42, "episode_budget": 50}.items():
            require_equal(planning.get(key), expected, f"{spec.key}.planning.{key}")
        outcomes[spec.key] = tuple(flags)

    assert shared_selection is not None and shared_action is not None
    return outcomes, shared_selection


def reconcile(outcomes: Mapping[str, Sequence[bool]], selection: Mapping[str, Any]) -> dict[str, Any]:
    baseline = outcomes["f_only"]
    cells: dict[str, Any] = {}
    for spec in CELL_SPECS:
        flags = outcomes[spec.key]
        rescues = [not old and new for old, new in zip(baseline, flags)]
        lost = [old and not new for old, new in zip(baseline, flags)]
        both_fail = [not old and not new for old, new in zip(baseline, flags)]
        union = [old or new for old, new in zip(baseline, flags)]
        cells[spec.key] = {
            "label": spec.label,
            "checkpoint": f"V1-{spec.variant.upper()} E{spec.epoch}",
            "training_loss": "L_C3" if spec.key == "c3" else "L_C",
            "score_path": SCORE_PATHS[spec.label],
            "success_count": sum(flags),
            "success_rate_percent": 2 * sum(flags),
            "historical_o50_count": spec.historical_o50_count,
            "historical_o50_rate_percent": 2 * spec.historical_o50_count,
            "delta_vs_f_count": sum(flags) - sum(baseline),
            "delta_vs_f_pp": 2 * (sum(flags) - sum(baseline)),
            "retained_f_successes": pair_ids(
                [old and new for old, new in zip(baseline, flags)]
            ),
            "new_rescues": pair_ids(rescues),
            "lost_f_successes": pair_ids(lost),
            "both_fail": pair_ids(both_fail),
            "f_or_method_count": sum(union),
            "f_or_method_rate_percent": 2 * sum(union),
        }
    any_success = [any(outcomes[key][index] for key in outcomes) for index in range(EPISODES)]
    any_rescue = [not baseline[index] and any_success[index] for index in range(EPISODES)]
    return {
        "schema_version": 1,
        "study": "V1-C and V1-C3 paired Cube O25 extension",
        "episode_count": EPISODES,
        "outcome_count": len(CELL_SPECS) * EPISODES,
        "selection_sha256": SELECTION_SHA256,
        "valid_row_ranks_sha256": RANKS_SHA256,
        "action_normalization_sha256": ACTION_SHA256,
        "v1_c_checkpoint_sha256": V1_C_CHECKPOINT_SHA256,
        "v1_c3_checkpoint_sha256": V1_C3_CHECKPOINT_SHA256,
        "f_success_pair_ids": pair_ids(baseline),
        "f_failure_pair_ids": pair_ids([not value for value in baseline]),
        "cells": cells,
        "oracle": {
            "f_plus_c3_count": cells["c3"]["f_or_method_count"],
            "f_plus_c3_rate_percent": cells["c3"]["f_or_method_rate_percent"],
            "f_plus_all_alternatives_count": sum(any_success),
            "f_plus_all_alternatives_rate_percent": 2 * sum(any_success),
            "all_alternative_rescue_pair_ids": pair_ids(any_rescue),
            "universal_failure_pair_ids": pair_ids([not value for value in any_success]),
            "interpretation": "Oracle unions use observed success labels and are not deployable gating results.",
        },
        "selection": {name: selection[name] for name in ("episode_indices", "start_steps", "goal_steps", "valid_row_ranks")},
    }


def build_paired_csv(outcomes: Mapping[str, Sequence[bool]], selection: Mapping[str, Any]) -> bytes:
    stream = io.StringIO(newline="")
    fieldnames = ["pair_id", "episode_index", "start_step", "goal_step", "valid_row_rank", *[spec.key for spec in CELL_SPECS], "f_or_c3", "any_method"]
    writer = csv.DictWriter(stream, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()
    for index in range(EPISODES):
        row = {
            "pair_id": f"P{index + 1:02d}",
            "episode_index": selection["episode_indices"][index],
            "start_step": selection["start_steps"][index],
            "goal_step": selection["goal_steps"][index],
            "valid_row_rank": selection["valid_row_ranks"][index],
        }
        for spec in CELL_SPECS:
            row[spec.key] = "1" if outcomes[spec.key][index] else "0"
        row["f_or_c3"] = "1" if outcomes["f_only"][index] or outcomes["c3"][index] else "0"
        row["any_method"] = "1" if any(outcomes[spec.key][index] for spec in CELL_SPECS) else "0"
        writer.writerow(row)
    return stream.getvalue().encode("utf-8")


def markdown_table(headers: Sequence[str], rows: Sequence[Sequence[str]]) -> list[str]:
    def clean(value: str) -> str:
        return value.replace("|", "\\|").replace("\n", "<br>")
    lines = ["| " + " | ".join(clean(value) for value in headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines.extend("| " + " | ".join(clean(str(value)) for value in row) + " |" for row in rows)
    return lines


def build_markdown_section(summary: Mapping[str, Any], outcomes: Mapping[str, Sequence[bool]], paired_sha256: str) -> str:
    cells = summary["cells"]
    selection = summary["selection"]
    lines = [
        MARKDOWN_START,
        "## O25 配对补测 V1 C 与 C3",
        "",
        "同一组 50 个 O25 start-goal pair 上，单方法最高为 **C3 State-V + First-Q2 alpha=.10：38/50 (76%)**；V1-C F-only 为 **37/50 (74%)**。前六行是先前补测的 V1-C E10 六种评分；最后一行是把此前 O50 得到 31/50 (62%) 的同一个 C3 scorer 原样移到 O25，未重新训练。C3 相对 F-only 新救回 6 个 pair，同时丢失 5 个原 F 成功 pair。若事后使用成功标签做 oracle 选择，F-only 与 C3 的并集为 **43/50 (86%)**；把全部六种替代评分也纳入，oracle 上限为 **44/50 (88%)**。这些 oracle 数字不是可部署结果。",
        "",
        "### 协议与证据指纹",
        "",
    ]
    lines += markdown_table(("字段", "锁定值"), (
        ("任务", "Cube O25；50 个固定 start-goal pair；goal step - start step = 25"),
        ("规划", "CEM；300 candidates；30 iterations；30 elites；planning seed 42；action block 5"),
        ("执行节奏", "除 G-only 外均 H=5/RH=5、每 25 环境步重规划；G-only 为 H=1/RH=1、每 5 步重规划"),
        ("V1-C checkpoint", f"E10/global step 127960；SHA-256 `{V1_C_CHECKPOINT_SHA256}`"),
        ("V1-C3 checkpoint", f"E12/global step 12000；SHA-256 `{V1_C3_CHECKPOINT_SHA256}`"),
        ("Selection / ranks", f"`{SELECTION_SHA256}` / `{RANKS_SHA256}`"),
        ("Action normalization", f"`{ACTION_SHA256}`"),
        ("逐-pair CSV", f"`reports/artifacts/actor_free_td_lewm_v1_c_c3_o25_20260906/paired_outcomes.csv`；SHA-256 `{paired_sha256}`"),
    ))
    lines += [
        "",
        "### 训练目标与推理评分定义",
        "",
        "V1-C 的六种 O25 评分共享同一个 E10 checkpoint 与训练目标 `L_C=mean(l)+mean_goal(q-qY)^2`；C3 使用 `L_C3=mean_i omega_tau(r_i)Huber_1(r_i)`。本轮只改变或复用推理评分，没有为 O25 重新训练。`Zcand` 表示在每次 CEM candidate population 内分别做 z-score。",
        "",
    ]
    score_rows = []
    for spec in CELL_SPECS:
        execution = "A1 后重规划；真实 z0" if spec.key == "g_only" else "A1-A5 后重规划"
        if spec.key == "mean_q":
            execution += "；q1 用真实 z0，q2-q5 用 F imagined states"
        elif spec.key == "c3":
            execution += "；EMA State-V 读 F imagined terminal，online G 读真实 z0"
        score_rows.append(
            (
                spec.label,
                SCORE_PATHS[spec.label],
                f"H{spec.horizon}/RH{spec.receding_horizon}",
                execution,
            )
        )
    lines += markdown_table(("评分", "CEM 最小化 cost", "H/RH", "执行与状态来源"), score_rows)
    lines += ["", "### 七种评分结果与 F-only 配对覆盖", "", "下表中‘救回’是 F-only 失败而该方法成功；‘丢失’是 F-only 成功而该方法失败；`F∪方法`是使用真实成功标签事后选择得到的 oracle，并非已实现的门控器。O50 仅作既有结果参照，不能与 O25 直接比较难度。", ""]
    result_rows = []
    for spec in CELL_SPECS:
        cell = cells[spec.key]
        result_rows.append((spec.label, cell["training_loss"], f"{cell['success_count']}/50 ({cell['success_rate_percent']}%)", f"{cell['historical_o50_count']}/50 ({cell['historical_o50_rate_percent']}%)", str(len(cell["retained_f_successes"])), str(len(cell["new_rescues"])), str(len(cell["lost_f_successes"])), str(len(cell["both_fail"])), f"{cell['delta_vs_f_pp']:+d} pp", f"{cell['f_or_method_count']}/50 ({cell['f_or_method_rate_percent']}%)"))
    lines += markdown_table(("评分", "训练 loss", "O25", "同 scorer 既有 O50", "保留 F 成功", "新救回", "丢失 F 成功", "两者均失败", "相对 F", "F∪方法 oracle"), result_rows)
    lines += ["", "### F-only 成功与失败 pair", "", f"- F-only 成功 37 个：{', '.join(summary['f_success_pair_ids'])}", f"- F-only 失败 13 个：{', '.join(summary['f_failure_pair_ids'])}", "", "### 每种方法相对 F-only 的逐-pair 转移", ""]
    transition_rows = []
    for spec in CELL_SPECS:
        cell = cells[spec.key]
        transition_rows.append((spec.label, ", ".join(cell["new_rescues"]) or "无", ", ".join(cell["lost_f_successes"]) or "无", ", ".join(cell["both_fail"]) or "无"))
    lines += markdown_table(("评分", "新救回 F 失败", "丢失 F 成功", "两者均失败"), transition_rows)
    lines += ["", "### P01 到 P50 完整结果", "", "`S` 表示成功，`F` 表示失败。", ""]
    pair_rows = []
    for index in range(EPISODES):
        pair_rows.append((f"P{index + 1:02d}", str(selection["episode_indices"][index]), str(selection["start_steps"][index]), str(selection["goal_steps"][index]), str(selection["valid_row_ranks"][index]), *[("S" if outcomes[spec.key][index] else "F") for spec in CELL_SPECS]))
    lines += markdown_table(("Pair", "Episode", "Start", "Goal", "Rank", "F", "G", "F+G", "First-Q", "Mean-Q", "First-Q2", "C3 combo"), pair_rows)
    oracle = summary["oracle"]
    lines += [
        "", "### 结论与下一步门控目标", "",
        f"- C3 是 O25 单方法最高值 38/50 (76%)，但它只是对 F-only 净增加 1 个成功：新增 {', '.join(cells['c3']['new_rescues'])}，同时丢失 {', '.join(cells['c3']['lost_f_successes'])}。",
        "- C3 与 F-only 的配对不一致数为 6 对 5，双侧 exact McNemar `p=1.0`；当前单 seed 不能据此声称统计显著优于 F。",
        f"- F-only + C3 的 oracle 并集是 43/50 (86%)。全部替代评分的新增覆盖并集为 {', '.join(oracle['all_alternative_rescue_pair_ids'])}，所以 F + 全部方法的 oracle 为 44/50 (88%)；其中 P25 只被 C3 救回，P32 只被 F+G tail 救回。",
        f"- 所有七种评分都失败的 pair 为 {', '.join(oracle['universal_failure_pair_ids'])}。",
        "- 下一步不是把 86% 或 88% 当成真实控制成绩，而是在与正式 O25 分离的开发 pairs 上训练或校准门控器：先判断保留 F 还是切换 C3，再考虑增加 F+G tail 的第三路，以覆盖 P32 类型。门控规则和阈值锁定后，才能在未见 test pairs 与多个 planning seeds 上报告可部署结果。",
        "- 口径限制：一个 training seed、一个 planning seed、同一正式 50 pairs；C3 的 alpha=.10 曾在同一 O50 上探索选择。O25 与 O50 的 offset、难度和执行协议不同，不可用百分比差直接宣称跨协议提升。",
        "", MARKDOWN_END,
    ]
    return "\n".join(lines)


def replace_markdown(base: bytes, section: str) -> bytes:
    text = base.decode("utf-8")
    start_count, end_count = text.count(MARKDOWN_START), text.count(MARKDOWN_END)
    if (start_count, end_count) == (0, 0):
        require_equal(sha256_bytes(base), BASE_MARKDOWN_SHA256, "base Markdown SHA-256")
        updated = text.rstrip() + "\n\n" + section.rstrip() + "\n"
    elif (start_count, end_count) == (1, 1):
        start = text.index(MARKDOWN_START)
        end = text.index(MARKDOWN_END, start) + len(MARKDOWN_END)
        if text[end:].strip():
            raise ValueError("O25 Markdown extension is not the final marked section")
        updated = text[:start].rstrip() + "\n\n" + section.rstrip() + "\n"
    else:
        raise ValueError("O25 Markdown markers are incomplete or duplicated")
    return updated.encode("utf-8")


def load_docx_helpers() -> ModuleType:
    path = REPOSITORY_ROOT / "scripts/build_results_td_v1.py"
    spec = importlib.util.spec_from_file_location("_results_td_v1_docx_helpers", path)
    if spec is None or spec.loader is None:
        raise ImportError(path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _set_table_borders(table: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for row in table.rows:
        for cell in row.cells:
            properties = cell._tc.get_or_add_tcPr()
            borders = properties.find(qn("w:tcBorders"))
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                properties.append(borders)
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                element = borders.find(qn(f"w:{edge}"))
                if element is None:
                    element = OxmlElement(f"w:{edge}")
                    borders.append(element)
                element.set(qn("w:val"), "single")
                element.set(qn("w:sz"), "4")
                element.set(qn("w:color"), "D9D9D9")


def _format_table(table: Any, helpers: ModuleType, *, font_size: float = 8.2, centered_from: int | None = None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    for cell in table.rows[0].cells:
        helpers._shade_cell(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string("000000")
            run.font.size = Pt(font_size)
            run.bold = True
    for row in table.rows[1:]:
        for index, cell in enumerate(row.cells):
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            if centered_from is not None and index >= centered_from:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(font_size)
    _set_table_borders(table)


def _remove_docx_extension(document: Any) -> bool:
    from docx.oxml.ns import qn
    starts = [p for p in document.paragraphs if p.text == DOCX_START]
    ends = [p for p in document.paragraphs if p.text == DOCX_END]
    if not starts and not ends:
        return False
    if len(starts) != 1 or len(ends) != 1:
        raise ValueError("DOCX O25 extension markers are incomplete or duplicated")
    body = document._element.body
    children = list(body)
    start_index = children.index(starts[0]._p)
    end_index = children.index(ends[0]._p)
    if end_index < start_index:
        raise ValueError("DOCX O25 markers are reversed")
    for child in children[end_index + 1 :]:
        if child.tag != qn("w:sectPr"):
            if child.tag != qn("w:p") or "".join(child.itertext()).strip():
                raise ValueError("DOCX O25 extension is not the final section")
    for child in children[start_index : end_index + 1]:
        body.remove(child)
    return True


def _set_running_matter(section: Any, helpers: ModuleType) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for header in (section.header, section.even_page_header, section.first_page_header):
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        helpers._set_run_font(paragraph.add_run("Results TD paired extension · Cube O25 · V1-C and V1-C3"), size=8.5, color="6B7280")
    for footer in (section.footer, section.even_page_footer, section.first_page_footer):
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        helpers._set_run_font(paragraph.add_run("Validated paired outcomes · Page "), size=8.5, color="6B7280")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        paragraph._p.append(field)


def build_docx(base_path: Path, summary: Mapping[str, Any], outcomes: Mapping[str, Sequence[bool]], paired_sha256: str) -> bytes:
    from docx import Document
    from docx.shared import Pt
    helpers = load_docx_helpers()
    document = Document(str(base_path))
    existed = _remove_docx_extension(document)
    if not existed:
        require_equal(file_sha256(base_path), BASE_DOCX_SHA256, "base DOCX SHA-256")
        helpers._configure_append_section(document)
    _set_running_matter(document.sections[-1], helpers)
    cells = summary["cells"]
    selection = summary["selection"]

    kicker = document.add_paragraph(style="Report Kicker")
    helpers._set_run_font(kicker.add_run(DOCX_START), size=9.5, color="5C6975", bold=True)
    title = document.add_heading("Cube O25 paired audit for V1-C and V1-C3", level=1)
    title.paragraph_format.page_break_before = False
    for run in title.runs:
        helpers._set_run_font(run, size=22, color="000000", bold=True)
    subtitle = document.add_paragraph()
    helpers._set_run_font(subtitle.add_run("Retention, rescue and regression on one fixed set of 50 start-goal pairs"), size=12, color="4B5563")
    intro = document.add_paragraph("On 50 fixed O25 pairs, F-only is 37/50 (74%) and the C3 combo is 38/50 (76%). The first six rows are the earlier V1-C E10 supplement; the final row reuses, without retraining, the exact C3 scorer that reached 31/50 (62%) on O50. C3 rescues six F failures but loses five F successes. Label-oracle unions are 43/50 (86%) for F+C3 and 44/50 (88%) across all scores; they are not deployable results.")
    for run in intro.runs:
        helpers._set_run_font(run, size=9.5, color="7A5A00", bold=True)

    document.add_heading("Protocol and evidence fingerprints", level=2)
    protocol_rows = (
        ("Task", "Cube O25; 50 fixed pairs; goal step minus start step = 25"),
        ("CEM", "300 candidates; 30 iterations; 30 elites; planning seed 42; action block 5"),
        ("Execution cadence", "All but G-only: H5/RH5 and replan every 25 environment steps. G-only: H1/RH1 and replan every 5 steps."),
        ("V1-C", f"E10 / step 127960 / {V1_C_CHECKPOINT_SHA256}"),
        ("V1-C3", f"E12 / step 12000 / {V1_C3_CHECKPOINT_SHA256}"),
        ("Selection", SELECTION_SHA256),
        ("Valid-row ranks", RANKS_SHA256),
        ("Action normalization", ACTION_SHA256),
        ("Paired CSV", f"{paired_sha256} · reports/artifacts/actor_free_td_lewm_v1_c_c3_o25_20260906/paired_outcomes.csv"),
    )
    table = helpers._add_table(document, headers=("Field", "Locked value"), rows=protocol_rows, widths=(3100, 11300))
    _format_table(table, helpers, font_size=8.4)

    document.add_heading("Training objectives and inference scores", level=2)
    document.add_paragraph(
        "V1-C shares E10 and L_C=mean(l)+mean_goal(q-qY)^2; C3 uses "
        "L_C3=mean_i omega_tau(r_i)Huber_1(r_i). O25 reuses both checkpoints. "
        "Zcand is the within-candidate-population z-score."
    )
    score_rows = []
    for spec in CELL_SPECS:
        execution = "Replan after A1; real z0" if spec.key == "g_only" else "Replan after A1-A5"
        if spec.key == "mean_q":
            execution += "; q1 uses real z0; q2-q5 use F-imagined states"
        elif spec.key == "c3":
            execution += "; EMA State-V reads the F terminal; online G reads real z0"
        score_rows.append(
            (
                spec.label,
                SCORE_PATHS[spec.label],
                f"H{spec.horizon}/RH{spec.receding_horizon}",
                execution,
            )
        )
    table = helpers._add_table(
        document,
        headers=("Score", "Cost minimized by CEM", "H/RH", "Execution and state source"),
        rows=score_rows,
        widths=(2800, 6000, 1400, 4200),
    )
    _format_table(table, helpers, font_size=7.8)

    document.add_heading("Seven scores and paired coverage relative to F-only", level=2)
    document.add_paragraph(
        "The first six V1-C E10 rows are the earlier O25 supplement. The final "
        "V1-C3 E12 row is this round's O25 test of the historical O50 62% scorer."
    )
    legend = document.add_paragraph("Yellow fill marks the highest standalone O25 result; blue fill marks the highest F OR method oracle. Retained means both F and the method succeed; rescue means F fails and the method succeeds; lost means F succeeds and the method fails.")
    for run in legend.runs:
        helpers._set_run_font(run, size=9.5, color="374151", bold=True)
    result_rows = []
    for spec in CELL_SPECS:
        cell = cells[spec.key]
        result_rows.append((f"V1-{spec.variant.upper()} E{spec.epoch}", spec.label, cell["training_loss"], f"{cell['success_count']}/50\n{cell['success_rate_percent']}%", f"{cell['historical_o50_count']}/50\n{cell['historical_o50_rate_percent']}%", str(len(cell["retained_f_successes"])), str(len(cell["new_rescues"])), str(len(cell["lost_f_successes"])), str(len(cell["both_fail"])), f"{cell['delta_vs_f_pp']:+d} pp", f"{cell['f_or_method_count']}/50\n{cell['f_or_method_rate_percent']}%"))
    table = helpers._add_table(document, headers=("Checkpoint", "Score", "Loss", "O25", "Prior O50", "Retained", "Rescued", "Lost", "Both fail", "Delta F", "F OR method"), rows=result_rows, widths=(1600, 3000, 800, 1100, 1100, 1000, 900, 900, 1000, 900, 2100))
    _format_table(table, helpers, font_size=7.1, centered_from=2)
    max_success = max(cells[spec.key]["success_count"] for spec in CELL_SPECS)
    max_union = max(cells[spec.key]["f_or_method_count"] for spec in CELL_SPECS)
    for row, spec in zip(table.rows[1:], CELL_SPECS):
        if cells[spec.key]["success_count"] == max_success:
            helpers._shade_cell(row.cells[3], "FFF2CC")
        if cells[spec.key]["f_or_method_count"] == max_union:
            helpers._shade_cell(row.cells[10], "DDEBF7")

    document.add_heading("Exact F-only successes and failures", level=2)
    document.add_paragraph("F-only succeeds on 37 pairs: " + ", ".join(summary["f_success_pair_ids"]))
    document.add_paragraph("F-only fails on 13 pairs: " + ", ".join(summary["f_failure_pair_ids"]))

    document.add_heading("Exact pair transitions relative to F-only", level=2)
    transition_rows = []
    for spec in CELL_SPECS:
        if spec.key == "f_only":
            continue
        cell = cells[spec.key]
        transition_rows.append((spec.label, ", ".join(cell["new_rescues"]) or "None", ", ".join(cell["lost_f_successes"]) or "None", ", ".join(cell["both_fail"]) or "None"))
    table = helpers._add_table(document, headers=("Score", "Newly rescued F failures", "Lost F successes", "Both fail"), rows=transition_rows, widths=(2800, 3800, 3800, 4000))
    _format_table(table, helpers, font_size=8.0)

    heading = document.add_heading("Complete P01 to P50 outcome matrix", level=2)
    heading.paragraph_format.page_break_before = True
    document.add_paragraph("S means success and F means failure. Light green marks a rescued F-only failure; light red marks a lost F-only success; gray marks a pair failed by both F-only and the method.")
    pair_rows = []
    for index in range(EPISODES):
        pair_rows.append((f"P{index + 1:02d}", str(selection["episode_indices"][index]), str(selection["start_steps"][index]), str(selection["goal_steps"][index]), str(selection["valid_row_ranks"][index]), *[("S" if outcomes[spec.key][index] else "F") for spec in CELL_SPECS]))
    table = helpers._add_table(document, headers=("Pair", "Episode", "Start", "Goal", "Rank", "F", "G", "F+G", "First-Q", "Mean-Q", "First-Q2", "C3 combo"), rows=pair_rows, widths=(650, 1200, 900, 900, 1650, 1300, 1300, 1300, 1300, 1300, 1300, 1300))
    _format_table(table, helpers, font_size=7.3, centered_from=0)
    baseline = outcomes["f_only"]
    for pair_index, row in enumerate(table.rows[1:]):
        for spec_index, spec in enumerate(CELL_SPECS[1:], start=6):
            old, new = baseline[pair_index], outcomes[spec.key][pair_index]
            if not old and new:
                helpers._shade_cell(row.cells[spec_index], "E2F0D9")
            elif old and not new:
                helpers._shade_cell(row.cells[spec_index], "FCE4D6")
            elif not old and not new:
                helpers._shade_cell(row.cells[spec_index], "E7E6E6")

    document.add_heading("Conclusions and the next gating objective", level=2)
    oracle = summary["oracle"]
    conclusions = (
        f"C3 is the best standalone O25 score at 38/50 (76%), a net gain of only one success over F-only. It rescues {', '.join(cells['c3']['new_rescues'])}, but loses {', '.join(cells['c3']['lost_f_successes'])}.",
        "C3 versus F-only has six new-only and five reference-only successes; the two-sided exact McNemar p-value is 1.0. This single-seed result does not establish statistically significant superiority over F.",
        f"The F+C3 oracle is 43/50 (86%). The oracle across F and all alternatives is 44/50 (88%), with rescue union {', '.join(oracle['all_alternative_rescue_pair_ids'])}. P25 is unique to C3 and P32 is unique to F+G tail. Every score fails on {', '.join(oracle['universal_failure_pair_ids'])}.",
        "Next, train or calibrate a gate on separate development pairs: first choose between F and C3, then test whether an F+G-tail third route covers P32-like cases. Lock the rule before evaluation on unseen test pairs and multiple planning seeds. Formal O25 success labels must not be used to implement the oracle.",
        "Boundary: one training seed and one planning seed. C3 alpha=.10 was selected exploratorily on the same O50 set. O25 and O50 differ in offset, difficulty and execution protocol, so percentage differences are not cross-protocol improvements.",
    )
    for text in conclusions:
        paragraph = document.add_paragraph(text)
        paragraph.style = document.styles["List Bullet"]
    end = document.add_paragraph()
    run = end.add_run(DOCX_END)
    run.font.hidden = True
    run.font.size = Pt(1)
    stream = io.BytesIO()
    document.save(stream)
    payload = stream.getvalue()
    if not payload.startswith(b"PK"):
        raise RuntimeError("python-docx did not produce OOXML")
    return payload


def atomic_write_many(outputs: Mapping[Path, bytes]) -> None:
    staged: dict[Path, Path] = {}
    try:
        for path, payload in outputs.items():
            path.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(mode="wb", dir=path.parent, prefix=f".{path.name}.", suffix=".tmp", delete=False) as stream:
                stream.write(payload)
                stream.flush()
                os.fsync(stream.fileno())
                staged[path] = Path(stream.name)
        for path, temporary in staged.items():
            os.replace(temporary, path)
    finally:
        for temporary in staged.values():
            temporary.unlink(missing_ok=True)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Append the validated V1-C/V1-C3 O25 paired extension to Results TD")
    parser.add_argument("--artifact-root", type=Path, default=ARTIFACT_ROOT)
    parser.add_argument("--markdown", type=Path, default=MARKDOWN_PATH)
    parser.add_argument("--docx", type=Path, default=DOCX_PATH)
    parser.add_argument("--project-docx", type=Path, default=PROJECT_DOCX_PATH)
    parser.add_argument("--validate-only", action="store_true")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    artifact_root = args.artifact_root.expanduser().resolve()
    outcomes, selection = validate_evidence(artifact_root)
    summary = reconcile(outcomes, selection)
    paired_payload = build_paired_csv(outcomes, selection)
    paired_sha = sha256_bytes(paired_payload)
    if args.validate_only:
        print("PASS: 7/7 O25 cells, 350/350 Boolean outcomes, one shared 50-pair selection")
        print(f"F-only=37/50; C3=38/50; F|C3=43/50; F|all=44/50; paired_csv_sha256={paired_sha}")
        return 0

    markdown_path = args.markdown.expanduser().resolve()
    docx_path = args.docx.expanduser().resolve()
    project_docx_path = args.project_docx.expanduser().resolve()
    if not markdown_path.is_file() or not docx_path.is_file() or not project_docx_path.is_file():
        raise FileNotFoundError("canonical Markdown and both DOCX inputs must exist")
    if file_sha256(docx_path) != file_sha256(project_docx_path):
        raise ValueError("repo and project Results TD DOCX inputs must be byte-identical")
    markdown_payload = replace_markdown(markdown_path.read_bytes(), build_markdown_section(summary, outcomes, paired_sha))
    docx_payload = build_docx(docx_path, summary, outcomes, paired_sha)
    summary["source_files_sha256"] = {str(spec.directory / "results.json"): spec.results_sha256 for spec in CELL_SPECS}
    summary["outputs"] = {
        "paired_outcomes_csv": str(PAIRED_PATH.relative_to(REPOSITORY_ROOT)),
        "paired_outcomes_sha256": paired_sha,
        "markdown": str(markdown_path),
        "markdown_sha256": sha256_bytes(markdown_payload),
        "docx": str(docx_path),
        "project_docx": str(project_docx_path),
        "docx_sha256": sha256_bytes(docx_payload),
    }
    reconciliation_payload = (json.dumps(summary, indent=2, sort_keys=True, ensure_ascii=False) + "\n").encode("utf-8")
    atomic_write_many({PAIRED_PATH: paired_payload, RECONCILIATION_PATH: reconciliation_payload, markdown_path: markdown_payload, docx_path: docx_payload, project_docx_path: docx_payload})
    print("PASS: 7/7 O25 cells, 350/350 Boolean outcomes, one shared 50-pair selection")
    print("PASS: F-only=37/50; C3=38/50; F|C3=43/50; F|all=44/50")
    print(f"Wrote {PAIRED_PATH}")
    print(f"Wrote {RECONCILIATION_PATH}")
    print(f"Wrote {markdown_path}")
    print(f"Wrote byte-identical DOCX copies: {docx_path} and {project_docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
