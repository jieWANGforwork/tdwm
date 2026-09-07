from __future__ import annotations

from pathlib import Path

import pytest

import tdwm.results.actor_free_td_lewm_o25_o100_master as master_module
from tdwm.results.actor_free_td_lewm_o25_o100_master import (
    DOCX_O25_MARKER,
    DOCX_O100_MARKER,
    SCORE_KEYS,
    SECTION_END,
    SECTION_START,
    MasterEvidence,
    ProtocolMasterError,
    load_master_evidence,
    update_docx_document,
    update_markdown_text,
)

REPOSITORY_ROOT = Path(__file__).resolve().parents[2]

EXPECTED = {
    "o25": (
        (
            "V1-C E10 / OSMesa*",
            "L_C",
            (37, 29, 35, 36, 30, 35, None),
        ),
        (
            "V1-C3 E12 / OSMesa*",
            "L_C3",
            (None, None, None, None, None, None, 38),
        ),
        (
            "V1-C4 objective-v1 E10 / EGL",
            "L_C4=L_vector+L_goal",
            (37, 32, 38, 32, 31, 34, None),
        ),
    ),
    "o100": (
        (
            "V1-C E10 / OSMesa",
            "L_C",
            (25, 24, 22, 32, 25, 26, None),
        ),
        (
            "V1-C3 E12 / OSMesa",
            "L_C3",
            (None, None, None, None, None, None, 25),
        ),
        (
            "V1-C4 objective-v1 E10 / EGL",
            "L_C4=L_vector+L_goal",
            (25, 25, 24, 27, 28, 28, None),
        ),
    ),
}

DOCX_HEADERS = (
    "Method / checkpoint / runtime",
    "Training loss",
    "F-only",
    "G/C4-only",
    "F+G/F+C4 tail",
    "First-Q\nalpha=.25",
    "Mean-Q",
    "First-Q2\nalpha=.25",
    "C3 State-V+First-Q2\nalpha=.10",
)


@pytest.fixture(scope="module")
def evidence() -> MasterEvidence:
    return load_master_evidence(REPOSITORY_ROOT)


def _expected_rates(counts: tuple[int | None, ...]) -> list[str]:
    return ["—" if count is None else f"{count}/50 ({2 * count}%)" for count in counts]


def _plain_markdown_cell(value: str) -> str:
    if value.startswith("◆ "):
        value = value[2:]
    return value.replace("**", "")


def _markdown_rows(block: str) -> list[list[str]]:
    table_lines = [line for line in block.splitlines() if line.startswith("| ")]
    assert len(table_lines) == 5
    return [
        [cell.strip() for cell in line.strip().strip("|").split("|")]
        for line in table_lines[2:]
    ]


def _cell_fill(cell: object, qn: object) -> str:
    shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    assert shading is not None
    return (shading.get(qn("w:fill")) or "").upper()


def test_loads_exactly_thirteen_current_formal_cells_per_protocol_and_skips_v0(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    loaded_paths: list[Path] = []
    original = master_module._load_result

    def observed_load_result(
        path: Path, *, protocol: str, expected_mode: str
    ) -> master_module.ScoreCell:
        loaded_paths.append(path)
        return original(path, protocol=protocol, expected_mode=expected_mode)

    monkeypatch.setattr(master_module, "_load_result", observed_load_result)
    evidence = load_master_evidence(REPOSITORY_ROOT)

    assert len(loaded_paths) == 26
    assert all(
        "actor_free_td_lewm_v1_c4_20260907" not in path.parts for path in loaded_paths
    )
    assert (
        sum(
            "actor_free_td_lewm_v1_c4_objective1_20260907" in path.parts
            for path in loaded_paths
        )
        == 12
    )

    for protocol in ("o25", "o100"):
        result = getattr(evidence, protocol)
        assert result.protocol == protocol
        assert len(result.selection) == 50
        assert len(result.rows) == 3
        assert (
            sum(cell is not None for row in result.rows for cell in row.cells.values())
            == 13
        )

        for row, (label, loss, counts) in zip(result.rows, EXPECTED[protocol]):
            assert row.label == label
            assert row.training_loss == loss
            assert tuple(row.cells) == SCORE_KEYS
            assert (
                tuple(
                    None if row.cells[key] is None else row.cells[key].count
                    for key in SCORE_KEYS
                )
                == counts
            )
            for key, expected_count in zip(SCORE_KEYS, counts):
                cell = row.cells[key]
                if expected_count is None:
                    assert cell is None
                else:
                    assert cell is not None
                    assert len(cell.outcomes) == 50
                    assert sum(cell.outcomes) == expected_count

        assert all("objective-v0" not in row.label for row in result.rows)
        assert (
            tuple(result.rows[1].cells[key] for key in SCORE_KEYS[:-1]) == (None,) * 6
        )
        assert result.rows[0].cells["state_v_plus_first_q2"] is None
        assert result.rows[2].cells["state_v_plus_first_q2"] is None


def test_markdown_inserts_two_independent_exact_tables_and_is_idempotent(
    evidence: MasterEvidence,
) -> None:
    source = "\n".join(
        (
            "# Results TD",
            "",
            "模型均不训练 Actor。",
            "",
            "## O25 配对补测 V1 C 与 C3",
            "",
            "paired details",
            "",
            "## 训练 / validation loss 证据",
            "",
            "loss evidence",
        )
    )
    updated = update_markdown_text(source, evidence)

    assert updated.count(SECTION_START) == 1
    assert updated.count(SECTION_END) == 1
    assert updated.count("| 方法 / checkpoint / runtime |") == 2
    assert updated.count("### O25 当前正式结果总表") == 1
    assert updated.count("### O100 当前正式结果总表") == 1
    assert "## O25 V1 C 与 C3 配对明细" in updated

    section = updated.split(SECTION_START, 1)[1].split(SECTION_END, 1)[0]
    o25_block, o100_block = section.split("### O100 当前正式结果总表", 1)
    blocks = {"o25": o25_block, "o100": o100_block}
    parsed_rows: list[list[str]] = []
    for protocol, block in blocks.items():
        rows = _markdown_rows(block)
        parsed_rows.extend(rows)
        for actual, (label, loss, counts) in zip(rows, EXPECTED[protocol]):
            assert actual[:2] == [label, f"`{loss}`"]
            assert [_plain_markdown_cell(value) for value in actual[2:]] == (
                _expected_rates(counts)
            )
        assert sum(value != "—" for row in rows for value in row[2:]) == 13

    assert all("objective-v0" not in row[0] for row in parsed_rows)
    assert sum("objective-v1" in row[0] for row in parsed_rows) == 2

    with pytest.raises(ProtocolMasterError, match="already contains"):
        update_markdown_text(updated, evidence)


@pytest.mark.parametrize("marker", (SECTION_START, SECTION_END))
def test_markdown_refuses_even_a_partial_existing_master_marker(
    evidence: MasterEvidence, marker: str
) -> None:
    with pytest.raises(ProtocolMasterError, match="already contains"):
        update_markdown_text(marker, evidence)


def test_docx_inserts_tables_and_locates_values_and_colors_by_header(
    evidence: MasterEvidence,
) -> None:
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn

    document = docx.Document()
    decoy = document.add_table(rows=1, cols=1)
    decoy.cell(0, 0).text = "unrelated table"
    for text in (
        "RESULTS TD / V1-C4 FORMAL EXTENSION END",
        "Cube O25 paired audit for V1-C and V1-C3",
        "Methods by inference score, relative to F-only",
        "Cube O100 paired evaluation: V1-C and V1-C3",
        "O100 method-by-inference-score matrix",
    ):
        document.add_paragraph(text)

    update_docx_document(document, evidence)

    tables = [
        table
        for table in document.tables
        if tuple(cell.text for cell in table.rows[0].cells) == DOCX_HEADERS
    ]
    assert len(tables) == 2
    by_protocol = {
        "o25" if table.rows[1].cells[2].text.startswith("37/50") else "o100": table
        for table in tables
    }
    assert set(by_protocol) == {"o25", "o100"}

    for protocol, table in by_protocol.items():
        expected = EXPECTED[protocol]
        assert len(table.rows) == 4
        assert all(len(row.cells) == 9 for row in table.rows)
        assert all(_cell_fill(cell, qn) == "17365D" for cell in table.rows[0].cells)

        count_matrix = [entry[2] for entry in expected]
        row_maxima = [
            max(count for count in counts if count is not None)
            for counts in count_matrix
        ]
        column_maxima = [
            max(counts[column] for counts in count_matrix if counts[column] is not None)
            for column in range(len(SCORE_KEYS))
        ]
        for row_index, (label, loss, counts) in enumerate(expected, start=1):
            row = table.rows[row_index]
            assert [cell.text for cell in row.cells[:2]] == [label, loss]
            assert [cell.text for cell in row.cells[2:]] == _expected_rates(counts)
            base_fill = "FFFFFF" if row_index % 2 == 1 else "F3F6FA"
            assert [_cell_fill(cell, qn) for cell in row.cells[:2]] == [
                base_fill,
                base_fill,
            ]
            for column, (cell, count) in enumerate(zip(row.cells[2:], counts)):
                row_best = count is not None and count == row_maxima[row_index - 1]
                column_best = count is not None and count == column_maxima[column]
                expected_fill = (
                    "B7DEE8"
                    if row_best and column_best
                    else "DDEBF7"
                    if column_best
                    else "FFF2CC"
                    if row_best
                    else base_fill
                )
                assert _cell_fill(cell, qn) == expected_fill

    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "Cube O25 paired audit for V1-C and V1-C3" not in texts
    assert "Cube O25 current formal results and paired details" in texts
    assert "Cube O100 paired evaluation: V1-C and V1-C3" not in texts
    assert "Cube O100 current formal results and paired details" in texts
    assert texts.count("V1 C and C3 paired details relative to F only") == 2
    assert texts.count(DOCX_O25_MARKER) == 1
    assert texts.count(DOCX_O100_MARKER) == 1

    body = document._element.body
    anchors = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "V1 C and C3 paired details relative to F only"
    ]
    ordered_tables = sorted(tables, key=lambda table: body.index(table._tbl))
    ordered_anchors = sorted(anchors, key=lambda paragraph: body.index(paragraph._p))
    assert all(
        body.index(table._tbl) < body.index(anchor._p)
        for table, anchor in zip(ordered_tables, ordered_anchors)
    )

    with pytest.raises(ProtocolMasterError, match="already contains"):
        update_docx_document(document, evidence)
