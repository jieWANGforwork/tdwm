from __future__ import annotations

import csv
import hashlib
import json
import math
import zlib
from copy import deepcopy
from pathlib import Path

import pytest

from tdwm.results.actor_free_td_lewm_v1_c4 import (
    C4_ACTION_EFFECT,
    C4_JOINT_OBJECTIVE,
    C4_TIME_ALIGNMENT,
    HISTORICAL_V0_CHECKPOINT_SHA256,
    HISTORICAL_V0_DOCX_END_MARKER,
    HISTORICAL_V0_SECTION_END,
    HISTORICAL_V0_SECTION_START,
    LOSS_METRICS,
    OBJECTIVE_VERSION,
    PROTOCOLS,
    SCORE_MODES,
    C4ReportEvidence,
    C4ResultsUpdateError,
    LossSeries,
    _analysis_lines,
    load_loss_series,
    load_report_evidence,
    update_docx_document,
    update_markdown_text,
    validate_historical_v0_summary,
    validate_summary,
)

_PRE_C4_FIXTURE_DIR = (
    Path(__file__).resolve().parents[1] / "fixtures" / "actor_free_td_lewm_v1_c4"
)


def _paired(reference: list[bool], candidate: list[bool], *, f_plus: bool) -> dict:
    both_success = sum(left and right for left, right in zip(reference, candidate))
    new = sum((not left) and right for left, right in zip(reference, candidate))
    lost = sum(left and (not right) for left, right in zip(reference, candidate))
    discordant = new + lost
    exact_p = 1.0
    if discordant:
        lower = min(new, lost)
        exact_p = min(
            1.0,
            2.0
            * sum(math.comb(discordant, index) for index in range(lower + 1))
            / (2**discordant),
        )
    value = {
        "reference_successes": sum(reference),
        "candidate_successes": sum(candidate),
        "both_success": both_success,
        "new": new,
        "lost": lost,
        "both_failure": 50 - both_success - new - lost,
        "delta_successes": sum(candidate) - sum(reference),
        "delta_percentage_points": 2.0 * (sum(candidate) - sum(reference)),
        "exact_mcnemar_p_two_sided": exact_p,
        "new_episode_positions": [
            index + 1
            for index, (left, right) in enumerate(zip(reference, candidate))
            if not left and right
        ],
        "lost_episode_positions": [
            index + 1
            for index, (left, right) in enumerate(zip(reference, candidate))
            if left and not right
        ],
    }
    if f_plus:
        value["f_plus_new_successes"] = sum(reference) + new
    return value


def _outcome_vector(protocol_index: int, mode_index: int, method_index: int) -> list[bool]:
    threshold = 18 + protocol_index * 3 + mode_index + method_index
    return [((index * 7 + mode_index * 3) % 50) < threshold for index in range(50)]


def _summary(checkpoint_sha: str = "c" * 64) -> dict:
    protocols = {}
    episode_matrix = []
    nested = {}
    for protocol_index, protocol in enumerate(PROTOCOLS):
        methods = {}
        for method_index, method_key in enumerate(("c4", "v1_c")):
            scores = {}
            for mode_index, mode in enumerate(SCORE_MODES):
                if mode == "f_only":
                    outcomes = _outcome_vector(protocol_index, 0, 0)
                else:
                    outcomes = _outcome_vector(protocol_index, mode_index, method_index)
                nested[(protocol, method_key, mode)] = outcomes
                method_sha = checkpoint_sha if method_key == "c4" else "d" * 64
                scores[mode] = {
                    "success_count": sum(outcomes),
                    "success_rate_percent": 2.0 * sum(outcomes),
                    "episode_successes": outcomes,
                    "source": {
                        "directory": f"/{protocol}/{method_key}/{mode}",
                        "results": {"path": "/results.json", "sha256": "1" * 64},
                        "protocol_manifest": {"path": "/manifest.json", "sha256": "2" * 64},
                        "episode_selection": {
                            "path": "/selection.json",
                            "sha256": "3" * 64,
                            "valid_row_ranks_sha256": "4" * 64,
                        },
                        "checkpoint": {"path": "/epoch_10.pt", "sha256": method_sha},
                    },
                }
            methods[method_key] = {
                "checkpoint_sha256": checkpoint_sha if method_key == "c4" else "d" * 64,
                "scores": scores,
            }
        f_reference = nested[(protocol, "c4", "f_only")]
        versus_f = {
            mode: _paired(f_reference, nested[(protocol, "c4", mode)], f_plus=True)
            for mode in SCORE_MODES
        }
        versus_c = {
            mode: _paired(
                nested[(protocol, "v1_c", mode)],
                nested[(protocol, "c4", mode)],
                f_plus=False,
            )
            for mode in SCORE_MODES
        }
        protocols[protocol] = {
            "selection": {
                "episode_indices": list(range(50)),
                "start_steps": list(range(50)),
                "goal_steps": [index + (25, 50, 100)[protocol_index] for index in range(50)],
                "valid_row_ranks": list(range(50)),
                "valid_row_ranks_sha256": "4" * 64,
            },
            "methods": methods,
            "comparisons": {
                "c4_vs_same_protocol_f_only": versus_f,
                "c4_vs_v1_c_same_score_mode": versus_c,
            },
        }
        for position in range(50):
            episode_matrix.append(
                {
                    "protocol": protocol,
                    "episode_position": position + 1,
                    "pair_id": f"{protocol.upper()}-P{position + 1:02d}",
                    "valid_row_rank": position,
                    "episode_index": position,
                    "start_step": position,
                    "goal_step": position + (25, 50, 100)[protocol_index],
                    "c4": {mode: nested[(protocol, "c4", mode)][position] for mode in SCORE_MODES},
                    "v1_c": {mode: nested[(protocol, "v1_c", mode)][position] for mode in SCORE_MODES},
                }
            )
    return {
        "schema_version": 1,
        "study": {
            "method": "actor_free_td_lewm_v1_c4",
            "objective_version": OBJECTIVE_VERSION,
            "training_objective": C4_JOINT_OBJECTIVE["objective"],
            "comparison_method": "actor_free_td_lewm_v1_c",
            "training_seed": 3072,
            "protocols": list(PROTOCOLS),
            "score_modes": list(SCORE_MODES),
            "episodes_per_protocol": 50,
            "paired_comparison": True,
            "c4_checkpoint_sha256": checkpoint_sha,
            "v1_c_checkpoint_sha256": "d" * 64,
        },
        "protocols": protocols,
        "episode_matrix": episode_matrix,
    }


def _summary_with_v1_c_f_only_backend_drift(protocol: str = "o50") -> dict:
    summary = _summary()
    c4_outcomes = list(
        summary["protocols"][protocol]["methods"]["c4"]["scores"]["f_only"][
            "episode_successes"
        ]
    )
    v1_c_outcomes = list(c4_outcomes)
    c4_success = next(index for index, success in enumerate(c4_outcomes) if success)
    c4_failure = next(index for index, success in enumerate(c4_outcomes) if not success)
    v1_c_outcomes[c4_success] = False
    v1_c_outcomes[c4_failure] = True

    score = summary["protocols"][protocol]["methods"]["v1_c"]["scores"]["f_only"]
    score["episode_successes"] = v1_c_outcomes
    score["success_count"] = sum(v1_c_outcomes)
    score["success_rate_percent"] = 2.0 * sum(v1_c_outcomes)
    summary["protocols"][protocol]["comparisons"][
        "c4_vs_v1_c_same_score_mode"
    ]["f_only"] = _paired(v1_c_outcomes, c4_outcomes, f_plus=False)
    for row in summary["episode_matrix"]:
        if row["protocol"] == protocol:
            row["v1_c"]["f_only"] = v1_c_outcomes[row["episode_position"] - 1]
    return summary


def _historical_v0_summary() -> dict:
    summary = _summary(HISTORICAL_V0_CHECKPOINT_SHA256)
    del summary["study"]["objective_version"]
    del summary["study"]["training_objective"]
    return summary


def _manifest() -> dict:
    return {
        "method": "actor_free_td_lewm_v1_c4",
        "variant": "c4",
        "objective_version": OBJECTIVE_VERSION,
        "seed": 3072,
        "protocol": {
            "method": "actor_free_td_lewm_v1_c4",
            "variant": "c4",
            "objective_version": OBJECTIVE_VERSION,
            "stage": "full_training",
            "seeds": [3072],
            "g": {
                "state_dim": 192,
                "task_dim": 192,
                "output_dim": 192,
                "action_input": "none",
                "action_effect": C4_ACTION_EFFECT,
                "successor_semantics": "includes_current_input_state",
                "actor": "none",
                "reward": "none",
            },
            "time_alignment": deepcopy(C4_TIME_ALIGNMENT),
            "joint_objective": deepcopy(C4_JOINT_OBJECTIVE),
            "training": {"epochs": 10, "optimizer_steps_per_epoch": 12_796},
        },
        "model": {
            "trainable_modules": ["online_g_c4"],
            "optimizer_scope": "exact_online_g_parameters_only",
            "trainable_lewm_parameters": 0,
        },
        "training": {
            "formal_optimizer_steps": 127_960,
            "configured_optimizer_steps": 127_960,
            "epochs": 10,
            "validation_skipped": False,
            "f_output_stop_gradient": True,
            "lewm_prediction_loss": False,
            "sigreg_loss": False,
            "loss_metrics": list(LOSS_METRICS),
        },
    }


def _write_metrics(path: Path) -> None:
    columns = ["epoch"] + [
        f"{stage}/{metric}{'_epoch' if stage == 'train' else ''}"
        for stage in ("train", "validation")
        for metric in LOSS_METRICS
    ]
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        for epoch in range(10):
            row = {"epoch": epoch}
            for stage_index, stage in enumerate(("train", "validation")):
                component_values = []
                for metric_index, metric in enumerate(LOSS_METRICS[:-1]):
                    value = float(10 + stage_index + metric_index + epoch)
                    component_values.append(value)
                    key = f"{stage}/{metric}{'_epoch' if stage == 'train' else ''}"
                    row[key] = value
                total_key = f"{stage}/c4_total_loss{'_epoch' if stage == 'train' else ''}"
                row[total_key] = sum(component_values)
            writer.writerow(row)


def _evidence(summary: dict) -> C4ReportEvidence:
    losses = {
        stage: {
            metric: LossSeries(tuple(float(index + offset) for index in range(1, 11)))
            for offset, metric in enumerate(LOSS_METRICS)
        }
        for stage in ("train", "validation")
    }
    return C4ReportEvidence(
        summary=summary,
        summary_path=Path("/tmp/summary.json"),
        summary_sha256="a" * 64,
        historical_v0_summary=_historical_v0_summary(),
        historical_v0_summary_path=Path("/tmp/historical_v0_summary.json"),
        historical_v0_summary_sha256="e" * 64,
        training_manifest=_manifest(),
        training_manifest_path=Path("/tmp/run_manifest.json"),
        training_manifest_sha256="b" * 64,
        metrics_path=Path("/tmp/metrics.csv"),
        metrics_sha256="c" * 64,
        losses=losses,
        checkpoint_path=Path("/tmp/epoch_10.pt"),
        checkpoint_sha256="c" * 64,
    )


def test_validate_summary_requires_exact_18_cell_900_outcome_matrix() -> None:
    value = validate_summary(_summary())
    assert len(value["protocols"]) == 3
    assert sum(
        len(value["protocols"][protocol]["methods"]["c4"]["scores"][mode]["episode_successes"])
        for protocol in PROTOCOLS
        for mode in SCORE_MODES
    ) == 900


def test_validate_summary_rejects_missing_cell_and_paired_drift() -> None:
    missing = _summary()
    del missing["protocols"]["o100"]["methods"]["c4"]["scores"]["g_only"]
    with pytest.raises(C4ResultsUpdateError, match="exactly six"):
        validate_summary(missing)

    drifted = _summary()
    drifted["protocols"]["o50"]["comparisons"]["c4_vs_same_protocol_f_only"]["g_only"]["new"] += 1
    with pytest.raises(C4ResultsUpdateError, match="must be recomputed"):
        validate_summary(drifted)


def test_validate_summary_accepts_backend_f_only_drift_but_revalidates_pairing() -> None:
    summary = _summary_with_v1_c_f_only_backend_drift()
    validated = validate_summary(summary)
    paired = validated["protocols"]["o50"]["comparisons"][
        "c4_vs_v1_c_same_score_mode"
    ]["f_only"]
    assert paired["new"] == 1
    assert paired["lost"] == 1

    drifted = deepcopy(summary)
    drifted["protocols"]["o50"]["comparisons"][
        "c4_vs_v1_c_same_score_mode"
    ]["f_only"]["new"] += 1
    with pytest.raises(C4ResultsUpdateError, match="must be recomputed"):
        validate_summary(drifted)


def test_validate_historical_v0_summary_requires_exact_legacy_identity_and_matrix() -> None:
    summary = _historical_v0_summary()
    before = deepcopy(summary)
    validated = validate_historical_v0_summary(summary)
    assert validated == before
    assert "objective_version" not in validated["study"]
    assert sum(
        len(
            validated["protocols"][protocol]["methods"]["c4"]["scores"][mode][
                "episode_successes"
            ]
        )
        for protocol in PROTOCOLS
        for mode in SCORE_MODES
    ) == 900

    wrong_checkpoint = _historical_v0_summary()
    wrong_checkpoint["study"]["c4_checkpoint_sha256"] = "f" * 64
    with pytest.raises(C4ResultsUpdateError, match="unexpected checkpoint"):
        validate_historical_v0_summary(wrong_checkpoint)

    relabelled = _historical_v0_summary()
    relabelled["study"]["objective_version"] = OBJECTIVE_VERSION
    with pytest.raises(C4ResultsUpdateError, match="pre-versioned study schema"):
        validate_historical_v0_summary(relabelled)

    drifted = _historical_v0_summary()
    drifted["protocols"]["o25"]["methods"]["c4"]["scores"]["g_only"][
        "success_count"
    ] += 1
    with pytest.raises(C4ResultsUpdateError, match="disagrees with its outcomes"):
        validate_historical_v0_summary(drifted)


def test_load_loss_series_requires_ten_complete_consistent_epochs(tmp_path: Path) -> None:
    metrics = tmp_path / "metrics.csv"
    _write_metrics(metrics)
    values = load_loss_series(metrics)
    assert values["train"]["c4_total_loss"].first == 21.0
    assert values["validation"]["c4_total_loss"].final == 41.0

    rows = list(csv.DictReader(metrics.read_text().splitlines()))
    rows[-1]["validation/goal_projection_loss"] = "nan"
    with metrics.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)
    with pytest.raises(C4ResultsUpdateError, match="NaN/Inf"):
        load_loss_series(metrics)


def test_load_report_evidence_binds_checkpoint_and_optional_png(tmp_path: Path) -> None:
    checkpoint = tmp_path / "epoch_10.pt"
    checkpoint.write_bytes(b"c4-checkpoint")
    checkpoint_sha = hashlib.sha256(checkpoint.read_bytes()).hexdigest()
    summary = tmp_path / "summary.json"
    summary.write_text(json.dumps(_summary(checkpoint_sha)), encoding="utf-8")
    historical_v0_summary = tmp_path / "historical_v0_summary.json"
    historical_v0_summary.write_text(
        json.dumps(_historical_v0_summary()),
        encoding="utf-8",
    )
    manifest = tmp_path / "run_manifest.json"
    manifest.write_text(json.dumps(_manifest()), encoding="utf-8")
    metrics = tmp_path / "metrics.csv"
    _write_metrics(metrics)
    def chunk(kind: bytes, payload: bytes) -> bytes:
        checksum = zlib.crc32(kind + payload) & 0xFFFFFFFF
        return len(payload).to_bytes(4, "big") + kind + payload + checksum.to_bytes(4, "big")

    png = tmp_path / "loss.png"
    png.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(
            b"IHDR",
            (1).to_bytes(4, "big")
            + (1).to_bytes(4, "big")
            + bytes((8, 2, 0, 0, 0)),
        )
        + chunk(b"IDAT", zlib.compress(bytes((0, 0, 0, 0))))
        + chunk(b"IEND", b"")
    )

    evidence = load_report_evidence(
        summary_path=summary,
        historical_v0_summary_path=historical_v0_summary,
        training_manifest_path=manifest,
        metrics_path=metrics,
        checkpoint_path=checkpoint,
        loss_plot_path=png,
    )
    assert evidence.checkpoint_sha256 == checkpoint_sha
    assert evidence.loss_plot_path == png
    assert evidence.historical_v0_summary_path == historical_v0_summary

    checkpoint.write_bytes(b"wrong")
    with pytest.raises(C4ResultsUpdateError, match="checkpoint bytes differ"):
        load_report_evidence(
            summary_path=summary,
            historical_v0_summary_path=historical_v0_summary,
            training_manifest_path=manifest,
            metrics_path=metrics,
            checkpoint_path=checkpoint,
        )


def test_markdown_update_keeps_one_master_table_and_adds_paired_c4_matrices() -> None:
    report = _PRE_C4_FIXTURE_DIR / "results_td_before_c4.md"
    summary = _summary_with_v1_c_f_only_backend_drift()
    evidence = _evidence(summary)
    updated = update_markdown_text(report.read_text(encoding="utf-8"), evidence)
    assert "## 27 个训练方法 × 7 种评分" in updated
    assert updated.count("| V1 | C4 |") == 1
    assert "| V1-C4 objective-v1 formal O50 | 1 | E10 |" in updated
    assert "| **TOTAL** | — | — | same locked O50 selection | **511** |" in updated
    assert updated.count("### Protocol by score matrix") == 1
    assert updated.count("### F-only reproducibility/backend audit") == 1
    assert updated.count("### Paired outcomes relative to same-protocol F-only") == 1
    assert updated.count(
        "### Paired outcomes relative to historical V1-C under the same score (descriptive)"
    ) == 1
    section = updated.split("<!-- RESULTS_TD_V1_C4_FORMAL_START -->", 1)[1]
    paired_lines = [line for line in section.splitlines() if line.startswith(("| O25 |", "| O50 |", "| O100 |"))]
    assert len(paired_lines) == 36  # 3 score rows + 3 audit rows + two 15-row tables
    backend_pair = summary["protocols"]["o50"]["comparisons"][
        "c4_vs_v1_c_same_score_mode"
    ]["f_only"]
    expected_audit_row = (
        f"| O50 | {backend_pair['reference_successes']}/50 "
        f"({2 * backend_pair['reference_successes']}%) | "
        f"{backend_pair['candidate_successes']}/50 "
        f"({2 * backend_pair['candidate_successes']}%) | "
        f"{backend_pair['new']} | {backend_pair['lost']} | "
        f"{backend_pair['delta_successes']:+d} | No |"
    )
    assert expected_audit_row in section
    assert "current C4 evaluation used EGL" in section
    assert "historical V1-C reference used OSMesa" in section
    assert "same-EGL V1-C F-only rechecks" in section
    assert "not an effect of the C4 G head" in section
    assert "Within-C4 comparisons" in section
    assert "descriptive rather than pure C4 method effects" in section
    paired_section = section.split(
        "### Paired outcomes relative to same-protocol F-only", 1
    )[1].split("### Training loss and evidence", 1)[0]
    assert "| O25 | F-only |" not in paired_section
    assert "18 C4 cells and 900 C4 Boolean outcomes" in updated
    assert updated.count(HISTORICAL_V0_SECTION_START) == 1
    assert updated.count(HISTORICAL_V0_SECTION_END) == 1
    assert updated.count("objective v0 historical record - superseded") == 1
    assert updated.count("objective v1 formal O25 O50 O100") == 1
    historical_section = updated.split(HISTORICAL_V0_SECTION_START, 1)[1].split(
        HISTORICAL_V0_SECTION_END,
        1,
    )[0]
    historical_summary = _historical_v0_summary()
    for protocol in PROTOCOLS:
        cells = []
        for mode in SCORE_MODES:
            count = sum(
                historical_summary["protocols"][protocol]["methods"]["c4"][
                    "scores"
                ][mode]["episode_successes"]
            )
            cells.append(f"{count}/50 ({2 * count}%)")
        expected = f"| {protocol.upper()} | " + " | ".join(cells) + " |"
        assert expected in historical_section
    assert "excluded from the current 511-cell O50 ledger" in historical_section
    assert "| V1 | C4 |" not in historical_section

    with pytest.raises(C4ResultsUpdateError, match="already contains"):
        update_markdown_text(updated, evidence)


def test_historical_v0_results_do_not_enter_main_ledger_or_winners() -> None:
    report = _PRE_C4_FIXTURE_DIR / "results_td_before_c4.md"
    evidence = _evidence(_summary())
    history = evidence.historical_v0_summary
    protocol = "o50"
    mode = "g_only"
    outcomes = [True] * 50
    score = history["protocols"][protocol]["methods"]["c4"]["scores"][mode]
    score["episode_successes"] = outcomes
    score["success_count"] = 50
    score["success_rate_percent"] = 100.0
    f_reference = history["protocols"][protocol]["methods"]["c4"]["scores"][
        "f_only"
    ]["episode_successes"]
    v1_c_reference = history["protocols"][protocol]["methods"]["v1_c"]["scores"][
        mode
    ]["episode_successes"]
    comparisons = history["protocols"][protocol]["comparisons"]
    comparisons["c4_vs_same_protocol_f_only"][mode] = _paired(
        f_reference,
        outcomes,
        f_plus=True,
    )
    comparisons["c4_vs_v1_c_same_score_mode"][mode] = _paired(
        v1_c_reference,
        outcomes,
        f_plus=False,
    )
    for row in history["episode_matrix"]:
        if row["protocol"] == protocol:
            row["c4"][mode] = True
    validate_historical_v0_summary(history)

    updated = update_markdown_text(report.read_text(encoding="utf-8"), evidence)
    historical_section = updated.split(HISTORICAL_V0_SECTION_START, 1)[1].split(
        HISTORICAL_V0_SECTION_END,
        1,
    )[0]
    master_row = next(
        line for line in updated.splitlines() if line.startswith("| V1 | C4 |")
    )
    winner_row = next(
        line for line in updated.splitlines() if line.startswith("| V1 fixed |")
    )
    assert "50/50 (100%)" in historical_section
    assert "50/50 (100%)" not in master_row
    assert "objective v0" not in winner_row
    assert "| **TOTAL** | — | — | same locked O50 selection | **511** |" in updated


def test_analysis_is_dynamic_evidence_driven_and_predeclares_next_steps() -> None:
    summary = _summary_with_v1_c_f_only_backend_drift()
    analysis = "\n".join(_analysis_lines(_evidence(summary)))

    for protocol in PROTOCOLS:
        assert f"{protocol.upper()} scorer pattern for state-only/action-through-F C4" in analysis
        assert f"{protocol.upper()} complementarity:" in analysis
    paired = summary["protocols"]["o50"]["comparisons"][
        "c4_vs_v1_c_same_score_mode"
    ]["f_only"]
    assert (
        f"O50 {paired['reference_successes']}/50 "
        f"({2 * paired['reference_successes']}%) -> "
        f"{paired['candidate_successes']}/50 "
        f"({2 * paired['candidate_successes']}%) "
        f"(New {paired['new']}, Lost {paired['lost']}, "
        f"delta {paired['delta_successes']:+d}/50)"
    ) in analysis
    assert "historical V1-C/OSMesa -> current C4/EGL" in analysis
    assert "Within-C4 comparisons" in analysis
    assert "cross rendering backends" in analysis
    assert "cannot isolate a causal effect" in analysis
    assert "F+New preserves F successes only by oracle construction" in analysis
    assert "train goal/vector 11/10 (1.10x)" in analysis
    assert "validation goal/vector 11/10 (1.10x)" in analysis
    assert "optimization scale, not usefulness of the goal signal" in analysis
    assert "First-Q2 minus First-Q is O25" in analysis
    assert "does not authorize choosing a scorer after seeing" in analysis
    assert "Next predeclared experiment 1:" in analysis
    assert "disjoint development split" in analysis
    assert "Next predeclared experiment 2:" in analysis
    assert "non-deployable F+New oracle ceiling" in analysis
    assert "multiple training seeds and planning seeds" in analysis
    assert "separately for O25, O50, and O100" in analysis
    assert "No scorer is selected post hoc" in analysis


def test_summary_validation_does_not_mutate_input() -> None:
    summary = _summary()
    before = deepcopy(summary)
    validate_summary(summary)
    assert summary == before


def test_markdown_recomputes_v1_fixed_markers_and_winner_row() -> None:
    report = _PRE_C4_FIXTURE_DIR / "results_td_before_c4.md"
    summary = _summary()
    summary["protocols"]["o50"]["methods"]["c4"]["scores"]["f_only"][
        "success_count"
    ] = 30
    updated = update_markdown_text(report.read_text(encoding="utf-8"), _evidence(summary))
    c_row = next(line for line in updated.splitlines() if line.startswith("| V1 | C |"))
    c4_row = next(line for line in updated.splitlines() if line.startswith("| V1 | C4 |"))
    winner_row = next(
        line for line in updated.splitlines() if line.startswith("| V1 fixed |")
    )
    assert c_row.split(" | ")[3].startswith("23/50")
    assert c4_row.split(" | ")[3].startswith("◆ **30/50")
    assert "| V1 fixed | C4 30/50 |" in winner_row
    assert (
        "所有固定 E10 单格的最高结果为 V1-C4 + F-only: 30/50 (60%)"
        in updated
    )
    assert "加入 C4 后固定评分中的最高单格为 **V1-C4 + F-only: 30/50 (60%)**" in updated
    assert "在原 477 格基础账的 24 个训练配置内" in updated


def test_docx_update_in_memory_has_one_c4_row_and_preserves_old_audit_hashes() -> None:
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn

    repository = Path(__file__).resolve().parents[2]
    source = _PRE_C4_FIXTURE_DIR / "results_td_before_c4.docx"
    document = docx.Document(source)
    summary = _summary_with_v1_c_f_only_backend_drift()

    update_docx_document(document, _evidence(summary), repository)

    assert len(document.tables[18].rows) == 29
    assert sum(row.cells[1].text == "C4" for row in document.tables[18].rows) == 1
    methods = [row.cells[1].text for row in document.tables[18].rows]
    assert methods[7:12] == ["C", "C2", "C3", "C4", "D"]
    c4_row = next(row for row in document.tables[18].rows if row.cells[1].text == "C4")
    assert c4_row.cells[0].text == "V1"
    assert c4_row.cells[8].text == "26/50 (52%)"
    assert c4_row.cells[9].text == "—"
    shading = c4_row.cells[8]._tc.get_or_add_tcPr().find(qn("w:shd"))
    assert shading is not None and shading.get(qn("w:fill")) == "B7DEE8"
    assert document.tables[24].rows[-1].cells[-1].text == "511"
    assert document.tables[39].rows[3].cells[1].text == (
        "0e5b541bdb11cf6d647fc1e679499a02c3aa430d64e37c8819d02c44e1dcb900"
    )
    assert sum(
        paragraph.text == "RESULTS TD / V1-C4 FORMAL EXTENSION END"
        for paragraph in document.paragraphs
    ) == 1
    assert sum(
        paragraph.text == HISTORICAL_V0_DOCX_END_MARKER
        for paragraph in document.paragraphs
    ) == 1
    assert len(document.tables) == 52
    historical_table = document.tables[46]
    assert len(historical_table.rows) == 4
    historical_summary = _historical_v0_summary()
    for protocol_index, protocol in enumerate(PROTOCOLS, start=1):
        assert historical_table.rows[protocol_index].cells[0].text == protocol.upper()
        for mode_index, mode in enumerate(SCORE_MODES, start=1):
            count = historical_summary["protocols"][protocol]["methods"]["c4"][
                "scores"
            ][mode]["success_count"]
            assert historical_table.rows[protocol_index].cells[mode_index].text == (
                f"{count}/50 ({2 * count}%)"
            )
    audit_table = document.tables[48]
    assert len(audit_table.rows) == 4
    assert audit_table.rows[0].cells[1].text == "V1-C / OSMesa"
    backend_pair = summary["protocols"]["o50"]["comparisons"][
        "c4_vs_v1_c_same_score_mode"
    ]["f_only"]
    assert audit_table.rows[2].cells[3].text == str(backend_pair["new"])
    assert audit_table.rows[2].cells[4].text == str(backend_pair["lost"])
    assert audit_table.rows[2].cells[6].text == "No"
    assert len(document.tables[49].rows) == 16
    assert all(row.cells[1].text != "F-only" for row in document.tables[49].rows[1:])
    assert len(document.tables[50].rows) == 16
    assert document.tables[50].rows[0].cells[2].text == "V1-C"
    assert len(document.sections) == 12
    historical_section = document.sections[-4]
    for header in (
        historical_section.header,
        historical_section.first_page_header,
        historical_section.even_page_header,
    ):
        assert header.is_linked_to_previous is False
        assert "objective v0 historical record" in header.paragraphs[0].text
    for footer in (
        historical_section.footer,
        historical_section.first_page_footer,
        historical_section.even_page_footer,
    ):
        assert footer.is_linked_to_previous is False
        assert "Superseded V1-C4 objective v0 evidence" in footer.paragraphs[0].text
    for c4_section in document.sections[-3:]:
        assert c4_section.different_first_page_header_footer is True
        for header in (
            c4_section.header,
            c4_section.first_page_header,
            c4_section.even_page_header,
        ):
            assert header.is_linked_to_previous is False
            assert "V1-C4 objective v1 formal O25 O50 O100" in header.paragraphs[0].text
        for footer in (
            c4_section.footer,
            c4_section.first_page_footer,
            c4_section.even_page_footer,
        ):
            assert footer.is_linked_to_previous is False
            assert "Validated V1-C4 objective v1 paired outcomes" in footer.paragraphs[0].text
    repeat = document.tables[51].rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader"))
    assert repeat is not None
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "w(g)=sqrt(192) z_g/||z_g||_2" in text
    assert "Every A_k is one 25D block of five consecutive 5D primitive actions" in text
    assert "tau=0.03, gamma=0.98, n<=50 primitive steps" in text
    assert "single state-only online branch" in text
    assert "x_i = stop-gradient F(z_i^real,a_i)" in text
    assert "L_vector + L_goal" in text
    assert "F-only reproducibility/backend audit" in text
    assert "current C4 evaluation used EGL" in text
    assert "historical V1-C reference used OSMesa" in text
    assert "same-EGL V1-C F-only rechecks" in text
    assert "not an effect of the C4 G head" in text
    assert "primary controlled comparisons" in text
    assert "descriptive rather than pure C4 method effects" in text
    assert "objective v0 historical record superseded" in text
    assert "excluded from the current 511-cell O50 ledger" in text
    assert "equal real and stopped-F-predicted" not in text
    assert "real z_i and stop-gradient F(z_(i-1),a_(i-1)) share the target" not in text

    with pytest.raises(C4ResultsUpdateError, match="already contains"):
        update_docx_document(document, _evidence(_summary()), repository)
